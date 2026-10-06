#!/usr/bin/env python3
"""Build docs/ORDER_TO_CASH.docx from correlation.json and the seed CSVs.

Requires python-docx. Run from the repository root:

    python3 docs/samples/o2c/build_use_case_docx.py
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOCS = ROOT.parents[1]
SEED = ROOT / "seed"
OUT = DOCS / "ORDER_TO_CASH.docx"

sys.path.insert(0, str(ROOT.parent / "c2p"))
from build_use_case_docx import (  # noqa: E402
    BLACK,
    add_field_table,
    add_para,
    add_table,
    nonempty_pairs,
    set_run_font,
)

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Inches  # noqa: E402


def load_csv(name: str) -> list[dict[str, str]]:
    with (SEED / name).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


def add_footer(section, text: str):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text + "  ·  ")
    set_run_font(run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    page_run = paragraph.add_run()
    set_run_font(page_run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    page_run._r.append(begin)
    page_run._r.append(instr)
    page_run._r.append(end)


def blank(value: str) -> str:
    return "" if value in ("", "00000000") else value


def main() -> None:
    c = json.loads((ROOT / "correlation.json").read_text(encoding="utf-8"))
    so = c["salesOrder"]
    cust = c["customer"]
    cash = c["cash"]
    profit = c["profit"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    add_footer(section, "Sales order 6321  ·  billing 0090005785  ·  receipt 1400000000")

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        style.font.size = Pt(size)
        style.font.bold = True

    core = doc.core_properties
    core.title = "Order to Cash — sales order 6321"
    core.subject = "ZEVO CDS extract use case with sample data"
    core.category = "Order to Cash"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_run_font(title.add_run("Order to Cash"), size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    set_run_font(subtitle.add_run("Use case with sample data"), size=14, color=RGBColor(0x1F, 0x4E, 0x79))
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    set_run_font(
        meta.add_run(
            "Reference document: sales order 6321 (stored as 0000006321)\n"
            "SAP S/4HANA CDS extracts · ZEVO_CDS_EXPLORER_2_FILE · run 20261006_154927\n"
            "Company code 1710 · sales organization 1710 · distribution channel 10 · USD"
        ),
        size=10,
        color=RGBColor(0x33, 0x33, 0x33),
    )

    doc.add_heading("1. Purpose", level=1)
    add_para(
        doc,
        "This use case shows how documents from a standard sales order are correlated "
        "through outbound delivery, goods issue, billing, the universal journal, and the "
        "incoming customer payment. The sample rows below are taken from the ZEVO file "
        "extracts for sales order 6321.",
    )
    add_para(
        doc,
        "ZEVO writes one CDS entity per file and one entity per ExtractCds call. It does "
        "not join documents. The client correlates the files with the keys in section 12. "
        "Full extracts are in docs/samples/o2c/source. The rows used here are in "
        "docs/samples/o2c/seed. Dates are YYYYMMDD, as stored in the CSV. Document numbers "
        "keep their leading zeros as extracted.",
    )

    doc.add_heading("2. Document flow", level=1)
    add_table(
        doc,
        ["Step", "Document", "What it carries for this case"],
        [
            ["1", "Sales order 0000006321", "Type TA, customer 0001000569, two bike items, net 14,800.00 USD"],
            ["2", "Delivery 0080006423", "First partial delivery: 15 Y240 and 7 Y200, picked and goods issued"],
            ["3", "Material document 4900008955", "Goods issue, movement type 601, accounting document 4900000004"],
            ["4", "Billing document 0090005785", "Type F2, net 3,240.00 USD, accounting document 9400000001"],
            ["5", "Journal 9400000001", "Debit receivable 12120000, credit revenue 41000000"],
            ["6", "Incoming payment 1400000000 (type DZ)", "3,240.00 USD on 20261006. Clears the receivable. Balance 0.00"],
            ["7", "Delivery 0080006426", "Second delivery for the remaining 61 Y240 and 15 Y200. Not yet picked, issued, or billed"],
        ],
    )
    add_para(
        doc,
        "The cash step is complete for the first delivery. The rest of the order is still "
        "in an open delivery and has no goods issue, billing, or journal line yet.",
    )

    doc.add_heading("3. Why this order", level=1)
    add_para(
        doc,
        "Sales order 6321 is the only order in the extract that runs from order entry "
        "through billing to a cleared customer receipt. The extract also contains order "
        "0000006320 (customer USCU_S01, billed on 0090005784, not paid) and service order "
        "0000006319 (type SRVO, no delivery or billing).",
    )

    doc.add_heading("4. Customer", level=1)
    add_para(
        doc,
        f"Sold-to, ship-to, and payer are all {cust['customer']}: {cust['bpFullName']} "
        f"({cust['country']} {cust['postalCode']}, {cust['street']}). Account group "
        f"{cust['accountGroup']}, created by {cust['createdByUser']} on {cust['creationDate']}.",
    )
    customer_row = load_csv("I_CUSTOMER.csv")[0]
    add_para(doc, "I_Customer", size=11, bold=True, space_after=4)
    add_field_table(
        doc,
        nonempty_pairs(
            customer_row,
            [
                "CUSTOMER",
                "CUSTOMERNAME",
                "CUSTOMERFULLNAME",
                "BPCUSTOMERFULLNAME",
                "CUSTOMERACCOUNTGROUP",
                "COUNTRY",
                "REGION",
                "POSTALCODE",
                "STREETNAME",
                "LANGUAGE",
                "CREATEDBYUSER",
                "CREATIONDATE",
            ],
        ),
    )

    doc.add_heading("5. Sales order 0000006321", level=1)
    add_para(
        doc,
        f"Standard order (type {so['salesOrderType']}), created by {so['createdByUser']} on "
        f"{so['creationDate']}. Customer purchase order {so['purchaseOrderByCustomer']} dated "
        f"{so['customerPurchaseOrderDate']}. Requested delivery {so['requestedDeliveryDate']}. "
        f"Payment terms {so['customerPaymentTerms']}, Incoterms {so['incoterms']}. "
        f"Overall delivery status {so['overallDeliveryStatus']} and process status "
        f"{so['overallProcessStatus']} on the header.",
    )
    so_row = load_csv("I_SALESORDER.csv")[0]
    add_para(doc, "I_SalesOrder — header", size=11, bold=True, space_after=4)
    add_field_table(
        doc,
        nonempty_pairs(
            so_row,
            [
                "SALESORDER",
                "SALESORDERTYPE",
                "CREATEDBYUSER",
                "CREATIONDATE",
                "SALESORDERDATE",
                "SALESORGANIZATION",
                "DISTRIBUTIONCHANNEL",
                "ORGANIZATIONDIVISION",
                "SOLDTOPARTY",
                "PURCHASEORDERBYCUSTOMER",
                "CUSTOMERPURCHASEORDERDATE",
                "REQUESTEDDELIVERYDATE",
                "TOTALNETAMOUNT",
                "TRANSACTIONCURRENCY",
                "CUSTOMERPAYMENTTERMS",
                "INCOTERMSCLASSIFICATION",
                "INCOTERMSLOCATION1",
                "BILLINGCOMPANYCODE",
                "CREDITCONTROLAREA",
                "SDPRICINGPROCEDURE",
                "OVERALLDELIVERYSTATUS",
                "OVERALLSDPROCESSSTATUS",
                "LASTCHANGEDATETIME",
            ],
        ),
    )
    add_para(doc, "I_SalesOrderItem — items", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Item", "Cat.", "Material", "Text", "Qty", "Confirmed", "Net price", "Net amount", "Plant", "SLoc", "Profit center", "Deliv. status"],
        [
            [
                it["salesOrderItem"],
                it["itemCategory"],
                it["material"],
                it["text"],
                it["orderQuantity"],
                it["confirmedQuantity"],
                it["netPrice"],
                it["netAmount"],
                it["plant"],
                it["storageLocation"],
                it["profitCenter"],
                it["deliveryStatus"],
            ]
            for it in so["items"]
        ],
        font=7,
    )
    add_para(
        doc,
        "12,160.00 + 2,640.00 = 14,800.00, the header TotalNetAmount. Both items are "
        "confirmed in full. Item 000020 has no storage location on the order; the delivery "
        "assigns 171A.",
    )

    doc.add_heading("6. Outbound deliveries", level=1)
    add_para(
        doc,
        "I_DeliveryDocumentItem points back to the order with ReferenceSDDocument and "
        "ReferenceSDDocumentItem (category C). Two deliveries exist for this order.",
    )
    add_para(doc, "I_DeliveryDocument — headers", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Delivery", "Type", "Created by", "Created", "Ship. point", "Ship-to", "Picking date", "Planned GI", "Actual GI", "Picking", "Goods mvmt", "Billing"],
        [
            [
                d["deliveryDocument"],
                d["deliveryDocumentType"],
                d["createdByUser"],
                d["creationDate"],
                d["shippingPoint"],
                d["shipToParty"],
                d["pickingDate"],
                d["plannedGoodsIssueDate"],
                blank(d["actualGoodsMovementDate"]),
                d["overallPickingStatus"],
                d["overallGoodsMovementStatus"],
                d["overallBillingStatus"],
            ]
            for d in c["deliveries"]
        ],
        font=7,
    )
    add_para(doc, "I_DeliveryDocumentItem — items", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Delivery", "Item", "Order item", "Material", "Qty", "Unit", "Plant", "SLoc", "Mvmt", "Picking", "Goods mvmt", "Billing"],
        [
            [
                d["deliveryDocument"],
                it["deliveryDocumentItem"],
                it["referenceSDDocumentItem"],
                it["material"],
                it["actualDeliveryQuantity"],
                it["unit"],
                it["plant"],
                it["storageLocation"],
                it["goodsMovementType"],
                it["pickingStatus"],
                it["goodsMovementStatus"],
                it["deliveryRelatedBillingStatus"],
            ]
            for d in c["deliveries"]
            for it in d["items"]
        ],
        font=7,
    )
    add_para(
        doc,
        "Status C is completed, A is not yet started. Delivery 0080006423 is picked, "
        "goods issued, and billed. Delivery 0080006426 was created by SAP_SYSTEM on "
        "20261006 for the remaining quantities and has no follow-on documents.",
    )

    doc.add_heading("7. Goods issue", level=1)
    add_para(
        doc,
        "I_GoodsMovementDocumentDEX carries DeliveryDocument and DeliveryDocumentItem. "
        "Movement type 601, inventory transaction type WL, goods-movement reference "
        "document type L (delivery). ReferenceDocument repeats the delivery number.",
    )
    add_table(
        doc,
        ["Material doc", "Year", "Item", "Delivery", "Dlv item", "Material", "Qty", "Unit", "Mvmt", "D/C", "Posting date", "Customer"],
        [
            [
                g["materialDocument"],
                g["materialDocumentYear"],
                g["materialDocumentItem"],
                g["deliveryDocument"],
                g["deliveryDocumentItem"],
                g["material"],
                g["quantity"],
                g["unit"],
                g["goodsMovementType"],
                g["debitCreditCode"],
                g["postingDate"],
                g["customer"],
            ]
            for g in c["goodsIssues"]
        ],
        font=7,
    )
    add_para(
        doc,
        "The goods issue posts accounting document 4900000004 (type WL): credit stock "
        "0013600000 and debit cost of goods sold 0054083000 at the material cost. "
        "Join with ReferenceDocumentType MKPF and ReferenceDocument = MaterialDocument. "
        "The COGS lines carry SalesDocument and SalesDocumentItem.",
    )

    doc.add_heading("8. Billing", level=1)
    bd = c["billingDocuments"][0]
    add_para(
        doc,
        f"Billing document {bd['billingDocument']} (type {bd['billingDocumentType']}, category "
        f"{bd['billingDocumentCategory']}), date {bd['billingDocumentDate']}, created by "
        f"{bd['createdByUser']}. Payer {bd['payerParty']}, payment terms "
        f"{bd['customerPaymentTerms']}. Net {bd['totalNetAmount']} {bd['currency']}, tax "
        f"{bd['totalTaxAmount']}. Accounting document {bd['accountingDocument']} in company "
        f"{bd['companyCode']} fiscal year {bd['fiscalYear']}; posting status "
        f"{bd['accountingPostingStatus']}.",
    )
    add_para(doc, "I_BillingDocument — header", size=11, bold=True, space_after=4)
    bd_row = load_csv("I_BILLINGDOCUMENT.csv")[0]
    add_field_table(
        doc,
        nonempty_pairs(
            bd_row,
            [
                "BILLINGDOCUMENT",
                "BILLINGDOCUMENTTYPE",
                "BILLINGDOCUMENTCATEGORY",
                "SDDOCUMENTCATEGORY",
                "CREATEDBYUSER",
                "CREATIONDATE",
                "BILLINGDOCUMENTDATE",
                "SALESORGANIZATION",
                "DISTRIBUTIONCHANNEL",
                "DIVISION",
                "SOLDTOPARTY",
                "PAYERPARTY",
                "CUSTOMERPAYMENTTERMS",
                "TOTALNETAMOUNT",
                "TOTALTAXAMOUNT",
                "TRANSACTIONCURRENCY",
                "COMPANYCODE",
                "FISCALYEAR",
                "ACCOUNTINGDOCUMENT",
                "ACCOUNTINGPOSTINGSTATUS",
                "ACCOUNTINGTRANSFERSTATUS",
                "DOCUMENTREFERENCEID",
            ],
        ),
    )
    add_para(doc, "I_BillingDocumentItem — items", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Item", "Ref. delivery", "Dlv item", "Sales order", "SO item", "Material", "Qty", "Net amount", "Cost", "Profit center", "Services rendered"],
        [
            [
                it["billingDocumentItem"],
                it["referenceSDDocument"],
                it["referenceSDDocumentItem"],
                it["salesDocument"],
                it["salesDocumentItem"],
                it["material"],
                it["billingQuantity"],
                it["netAmount"],
                it["costAmount"],
                it["profitCenter"],
                it["servicesRenderedDate"],
            ]
            for it in bd["items"]
        ],
        font=7,
    )
    add_para(
        doc,
        "The billing items reference the delivery (ReferenceSDDocumentCategory J) and also "
        "carry SalesDocument / SalesDocumentItem, so the billing item joins to the order "
        "directly. 2,400.00 + 840.00 = 3,240.00. The cost amounts (1,301.70 + 476.49 = "
        "1,778.19) equal the COGS posted by the goods issue.",
    )

    doc.add_heading("9. Universal journal", level=1)
    add_para(
        doc,
        "I_GLAccountLineItemRawData, company code 1710, fiscal year 2026, ledger 0L. "
        "The billing document is linked by AccountingDocument on I_BillingDocument, and "
        "also by ReferenceDocumentType VBRK with ReferenceDocument = BillingDocument. "
        "The goods issue is linked by ReferenceDocumentType MKPF. Revenue and COGS lines "
        "carry SalesDocument, SalesDocumentItem, SoldProduct, and Customer.",
    )
    add_para(doc, "G/L accounts used (I_GLAccount, chart YCOA, company 1710)", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["G/L account", "External", "Group", "Reconciliation", "P&L", "Role on this chain"],
        [
            ["0011002000", "11002000", "FIN.", "", "", "Debited by the incoming payment (type DZ)"],
            ["0012120000", "12120000", "ABST", "D", "", "Customer receivables (financial account type D)"],
            ["0013600000", "13600000", "MAT.", "", "", "Stock, credited at goods issue"],
            ["0041000000", "41000000", "ERG.", "", "X", "Revenue, credited by billing"],
            ["0054083000", "0054083000", "ERG.", "", "X", "Cost of goods sold, debited at goods issue"],
        ],
    )
    add_para(
        doc,
        "I_GLAccount has no account-name column. Roles come from the posting pattern and the "
        "account attributes.",
        space_after=6,
    )
    add_para(doc, "Journal lines for this sales order", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["FI document", "Type", "Line", "G/L account", "D/C", "Amount USD", "Qty", "SO item", "Ref. type", "Reference document", "Clearing doc", "Clearing date"],
        [
            [
                l["accountingDocument"],
                l["accountingDocumentType"],
                l["ledgerGLLineItem"],
                l["glAccount"],
                l["debitCreditCode"],
                l["amountInCompanyCodeCurrency"],
                blank(l["quantity"]) if l["quantity"] not in ("0.000",) else "",
                l["salesDocumentItem"] if l["salesDocument"] else "",
                l["referenceDocumentType"],
                l["referenceDocument"],
                l["clearingAccountingDocument"],
                blank(l["clearingDate"]),
            ]
            for l in c["journalLines"]
        ],
        font=7,
    )
    add_para(
        doc,
        "Document 4900000004 (WL) moves 1,778.19 from stock to cost of goods sold. "
        "Document 9400000001 (RV) debits the receivable 3,240.00 and credits revenue "
        "2,400.00 and 840.00. Document 1400000000 (DZ) debits 0011002000 and credits the "
        "receivable; its clearing document is itself. The receivable line on 9400000001 "
        "carries the same clearing document and date.",
    )

    doc.add_heading("10. Quantity and amount reconciliation", level=1)
    add_table(
        doc,
        ["SO item", "Material", "Ordered", "Order net", "Goods issued", "In open delivery", "Billed qty", "Billed net", "Not yet billed qty", "Not yet billed net"],
        [
            [
                r["salesOrderItem"],
                r["material"],
                r["orderQuantity"],
                r["orderNetAmount"],
                r["goodsIssuedQuantity"],
                r["inOpenDeliveryQuantity"],
                r["billedQuantity"],
                r["billedNetAmount"],
                r["notYetBilledQuantity"],
                r["notYetBilledNetAmount"],
            ]
            for r in c["quantityReconciliation"]
        ],
        font=7,
    )
    add_para(
        doc,
        f"Billed revenue {profit['revenue']} {profit['currency']}, cost of goods sold "
        f"{profit['costOfGoodsSold']}, gross margin {profit['grossMargin']} on the first "
        "delivery. Issued quantity equals billed quantity. 11,560.00 of the order remains "
        "to be delivered and billed.",
    )

    doc.add_heading("11. Cash", level=1)
    add_para(
        doc,
        f"Incoming payment {cash['accountingDocument']} (document type "
        f"{cash['accountingDocumentType']}), posting date {cash['postingDate']}, amount "
        f"{cash['receivedAmount']} {cash['currency']} from customer {cash['customer']}. "
        f"Receivable balance on {cash['receivableGlAccount']} after clearing: "
        f"{cash['receivableBalance']} {cash['currency']}.",
    )
    add_para(doc, "Payment lines (I_GLAccountLineItemRawData, ledger 0L)", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["FI document", "Line", "G/L account", "D/C", "Posting key", "Amount USD", "Customer", "Offsetting account", "Clearing doc", "Clearing date"],
        [
            [
                cash["accountingDocument"],
                l["ledgerGLLineItem"],
                l["glAccount"],
                l["debitCreditCode"],
                l["postingKey"],
                l["amountInCompanyCodeCurrency"],
                l["customer"],
                l["offsettingAccount"],
                l["clearingAccountingDocument"],
                blank(l["clearingDate"]),
            ]
            for l in cash["lines"]
        ],
    )
    add_para(doc, "Receivable cleared by this payment", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["FI document", "Billing document", "Amount USD", "Clearing doc", "Clearing date"],
        [
            [
                r["accountingDocument"],
                r["billingDocument"],
                r["amount"],
                r["clearingAccountingDocument"],
                r["clearingDate"],
            ]
            for r in cash["clearedReceivables"]
        ],
    )
    add_para(
        doc,
        "Debit 3,240.00 on G/L 0011002000, posting key 40, assignment 20261006. Credit "
        "3,240.00 on receivable 0012120000, posting key 15, customer 0001000569. "
        "Reference document type is BKPF. The payment has no SD document number; the link "
        "to the billing document is ClearingAccountingDocument on the RV receivable line.",
    )

    doc.add_heading("12. How the documents are correlated", level=1)
    add_table(
        doc,
        ["From", "To", "Keys"],
        [[j["from"], j["to"], "; ".join(j["keys"]) + (f". {j['note']}" if j.get("note") else "")] for j in c["joinKeys"]],
        font=8,
    )
    add_para(
        doc,
        "SD numbers and FI numbers differ. Billing document 0090005785 posts to accounting "
        "document 9400000001. Material document 4900008955 posts to accounting document "
        "4900000004. Join SD documents to the journal on ReferenceDocument and "
        "ReferenceDocumentType, or use AccountingDocument on I_BillingDocument.",
    )

    doc.add_heading("13. Extract calls for this reference", level=1)
    add_para(
        doc,
        "The files were written by ZEVO_CDS_EXPLORER_2_FILE in run 20261006_154927. The "
        "same chain over OData is one ExtractCds call per entity. Example for the order header:",
    )
    add_para(
        doc,
        "GET /sap/opu/odata/sap/ZEVO_CDS_EXTRACT_SRV/ExtractCds"
        "?EntityName='I_SalesOrder'"
        "&Filter='SalesOrder eq ''0000006321'''"
        "&Format='jsonrows'&Skip='0'&Top='10'&$format=json",
        size=9,
    )
    add_table(
        doc,
        ["Entity", "Filter"],
        [
            ["I_SalesOrderItem", "SalesOrder eq '0000006321'"],
            ["I_DeliveryDocumentItem", "ReferenceSDDocument eq '0000006321'"],
            ["I_DeliveryDocument", "DeliveryDocument eq '0080006423' or DeliveryDocument eq '0080006426'"],
            ["I_GoodsMovementDocumentDEX", "DeliveryDocument eq '0080006423'"],
            ["I_BillingDocumentItem", "SalesDocument eq '0000006321'"],
            ["I_BillingDocument", "BillingDocument eq '0090005785'"],
            ["I_GLAccountLineItemRawData", "CompanyCode eq '1710' and SourceLedger eq '0L' and (AccountingDocument eq '4900000004' or AccountingDocument eq '9400000001' or AccountingDocument eq '1400000000')"],
            ["I_Customer", "Customer eq '0001000569'"],
            ["I_GLAccount", "CompanyCode eq '1710' and GLAccount in 0011002000, 0012120000, 0013600000, 0041000000, 0054083000 (as an or chain)"],
        ],
        font=8,
    )
    add_para(
        doc,
        "Order numbers in the filter need the stored leading zeros. There is no $expand and "
        "no in operator; chain or conditions.",
    )

    doc.add_heading("14. Sample-data files", level=1)
    add_para(
        doc,
        "CSV delimiter is semicolon. Headers are the CDS element names in uppercase. "
        "Regenerate the seed and docs/samples/o2c/correlation.json with "
        "python3 docs/samples/o2c/build_seed.py, then regenerate this document with "
        "python3 docs/samples/o2c/build_use_case_docx.py.",
    )

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
