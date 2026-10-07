#!/usr/bin/env python3
"""Build docs/REQUISITION_TO_PAYMENT.docx in the same 13 chapters as the other use cases.

    python3 docs/samples/rtp/build_use_case_docx.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parents[1] / "REQUISITION_TO_PAYMENT.docx"

sys.path.insert(0, str(ROOT.parent / "c2p"))
from build_use_case_docx import BLACK, add_para, add_table, set_run_font  # noqa: E402


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Requisition 10001634  ·  contract 4600000042  ·  PO 4500002148  ·  ")
    set_run_font(run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    page_run = paragraph.add_run()
    set_run_font(page_run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    page_run._r.append(begin)
    page_run._r.append(instr)
    page_run._r.append(end)


def main() -> None:
    c = json.loads((ROOT / "correlation.json").read_text(encoding="utf-8"))
    contract = c["contract"]
    po = c["purchaseOrder"]
    invoice = c["supplierInvoice"]
    seed = c["seed"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    add_footer(section)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        style.font.size = Pt(size)
        style.font.bold = True
    doc.core_properties.title = "Requisition to Payment — 10001634 / 4600000042 / 4500002148"
    doc.core_properties.subject = "ZEVO CDS extract use case with sample data"
    doc.core_properties.category = "Requisition to Payment"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_run_font(title.add_run("Requisition to Payment"), size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    set_run_font(subtitle.add_run("Use case with sample data"), size=14, color=RGBColor(0x1F, 0x4E, 0x79))
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    set_run_font(
        meta.add_run(
            "Reference documents: requisition 10001634, contract 4600000042, purchase order 4500002148\n"
            "SAP S/4HANA DEX extracts · ZEVO_CDS_EXPLORER_2_FILE · run 20261007_142521\n"
            "Company code 1710 · purchasing organization 1710 · USD"
        ),
        size=10,
        color=RGBColor(0x33, 0x33, 0x33),
    )

    doc.add_heading("1. Purpose", level=1)
    add_para(
        doc,
        "This use case follows a released purchase requisition into a contract, then into "
        "release purchase order 4500002148, the goods receipt, and the supplier invoice. "
        "The sample rows are the ZEVO extracts for that chain.",
    )
    add_para(
        doc,
        "ZEVO writes one CDS entity per file. The client correlates the files with the keys "
        "in section 11. The requisition item does not store the contract number. The "
        "purchase-order item and the contract history do. Dates are YYYYMMDD.",
    )

    doc.add_heading("2. Document flow", level=1)
    add_table(
        doc,
        ["Step", "Document", "What it carries"],
        [
            ["1", "Requisition 0010001634", "Printer paper and toner, processing status K, account assignment K"],
            ["2", "Contract 4600000042", "Type CWK, same items and quantities, lower prices, validity through 20271231"],
            ["3", "Purchase order 4500002148", "Release of 20 paper and 5 toner, total 219.00 USD"],
            ["4", "Material document 5000002952", "Goods receipt of the full release quantity, movement 101"],
            ["5", "Supplier invoice 5100001601", "SUPP.INV.0004, header gross 221.00, item amounts 219.00"],
            ["6", "Payment", "Not in this extract. There is no universal-journal file."],
        ],
    )

    doc.add_heading("3. Business partner", level=1)
    add_para(
        doc,
        f"Supplier {seed['supplier']} is {seed['supplierName']}. The same number is the "
        "business partner, the supplier on the requisition, the contract, and the purchase "
        "order, and the invoicing party on the invoice.",
    )

    doc.add_heading("4. Source document", level=1)
    first = c["requisitionItems"][0]
    add_para(
        doc,
        f"Purchase requisition 0010001634, type {first['requisitionType']}, company code "
        f"{first['companyCode']}, plant {first['plant']}, purchasing group {first['purchasingGroup']}. "
        f"Created by {first['createdByUser']} on {first['creationDate']}. Release status "
        f"{first['releaseStatus']}, release date {first['releaseDate']}. Processing status K "
        "means the item was converted to a contract.",
    )
    add_para(doc, "C_PurchaseRequisitionItemDEX", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Item", "Text", "Qty", "Price", "Net amount", "Delivery", "Supplier", "Acct", "Status"],
        [
            [
                row["item"],
                row["text"],
                row["quantity"],
                row["price"],
                row["netAmount"],
                row["deliveryDate"],
                row["supplier"],
                row["accountAssignmentCategory"],
                row["processingStatus"],
            ]
            for row in c["requisitionItems"]
        ],
        font=8,
    )
    add_para(
        doc,
        "5,000.00 + 2,500.00 = 7,500.00 at the requisition prices. PurchaseContract on both "
        "items is blank.",
    )

    doc.add_heading("5. Follow-on document", level=1)
    add_para(
        doc,
        f"Contract {contract['purchaseContract']}, type {contract['purchaseContractType']}, "
        f"created by {contract['createdByUser']} on {contract['creationDate']}. Validity "
        f"{contract['validityStartDate']} to {contract['validityEndDate']}. Payment terms "
        f"{contract['paymentTerms']}. The item texts, target quantities, supplier, plant, "
        "and account assignment category match the requisition. The contract prices are lower.",
    )
    add_para(doc, "C_PurchaseContractItemDEX", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Item", "Text", "Target qty", "Net price", "Target amount", "Plant", "Acct"],
        [
            [
                row["item"],
                row["text"],
                row["targetQuantity"],
                row["netPrice"],
                row["targetAmount"],
                row["plant"],
                row["accountAssignmentCategory"],
            ]
            for row in contract["items"]
        ],
    )
    add_para(
        doc,
        f"Release purchase order {po['purchaseOrder']}, type {po['purchaseOrderType']}, "
        f"date {po['purchaseOrderDate']}, created by {po['createdByUser']}. Header total "
        f"{po['totalAmount']} {po['documentCurrency']}. Each item stores the contract and "
        "the contract item. The release uses the contract price.",
    )
    add_para(doc, "C_PurchaseOrderItemDEX", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["PO item", "Contract", "Contract item", "Text", "Qty", "Net price", "Net amount", "Acct"],
        [
            [
                row["item"],
                row["purchaseContract"],
                row["purchaseContractItem"],
                row["text"],
                row["quantity"],
                row["netPrice"],
                row["netAmount"],
                row["accountAssignmentCategory"],
            ]
            for row in po["items"]
        ],
        font=8,
    )
    add_para(doc, "C_PurchaseContractHistoryDEX", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Contract item", "Release order", "Release item", "Qty", "Amount"],
        [
            ["00010", "4500002148", "00010", "20.000", "98.00"],
            ["00020", "4500002148", "00020", "5.000", "121.00"],
        ],
    )

    doc.add_heading("6. Goods movement", level=1)
    add_para(
        doc,
        "History document type 1, category E, movement type 101. Both items are on material "
        "document 5000002952, posted 20261009. The received quantity equals the released quantity.",
    )
    add_table(
        doc,
        ["Material document", "Year", "Item", "PO item", "Qty", "Movement", "Posting date"],
        [
            [
                row["materialDocument"],
                row["year"],
                row["item"],
                row["purchaseOrderItem"],
                row["quantity"],
                row["goodsMovementType"],
                row["postingDate"],
            ]
            for row in c["materialDocuments"]
        ],
    )

    doc.add_heading("7. Invoice", level=1)
    add_para(
        doc,
        f"Supplier invoice {invoice['supplierInvoice']}, party reference {invoice['reference']}, "
        f"posting date {invoice['postingDate']}, invoicing party {invoice['invoicingParty']}, "
        f"status {invoice['status']}. Header gross {invoice['grossAmount']} {invoice['currency']}. "
        "The two item amounts are 98.00 and 121.00, together 219.00.",
    )
    add_table(
        doc,
        ["Invoice", "Item", "PO item", "Qty", "Amount"],
        [
            [
                invoice["supplierInvoice"],
                row["item"],
                row["purchaseOrderItem"],
                row["quantity"],
                row["amount"],
            ]
            for row in invoice["items"]
        ],
    )

    doc.add_heading("8. Accounting", level=1)
    gl = c["glAccount"]
    add_para(
        doc,
        f"Account assignment category K. Both purchase-order items are assigned to cost center "
        f"{gl['costCenter']}, profit center YB600, and G/L {gl['glAccount']} "
        f"(external {gl['external']}, group {gl['group']}, profit-and-loss account). "
        "This extract set has no universal-journal file, so the accounting document for the "
        "invoice is not in the sample.",
    )
    add_table(
        doc,
        ["PO item", "Assignment", "Cost center", "G/L account", "Quantity", "Profit center"],
        [
            [
                row["purchaseOrderItem"],
                row["accountAssignmentNumber"],
                row["costCenter"],
                row["glAccount"],
                row["quantity"],
                row["profitCenter"],
            ]
            for row in c["accountAssignments"]
        ],
    )

    doc.add_heading("9. Quantity and amount reconciliation", level=1)
    add_table(
        doc,
        ["Text", "Req. qty", "Req. price", "Contract qty", "Contract price", "Released", "Received", "Invoiced", "Not released"],
        [
            [
                row["text"],
                row["requisitionQuantity"],
                row["requisitionPrice"],
                row["contractTargetQuantity"],
                row["contractPrice"],
                row["releasedQuantity"],
                row["receivedQuantity"],
                row["invoicedQuantity"],
                row["unreleasedQuantity"],
            ]
            for row in c["quantityReconciliation"]
        ],
        font=7,
    )
    add_para(
        doc,
        "The release order is fully received and fully invoiced. The contract still has "
        "980 of paper and 95 of toner not released. Paper moved from requisition price 5.00 "
        "to contract price 4.90. Toner moved from 25.00 to 24.20.",
    )

    doc.add_heading("10. Clearing", level=1)
    add_para(
        doc,
        "Supplier invoice 5100001601 is posted. A payment would be a later journal document "
        "whose clearing fields point at the vendor line of that invoice. This extract set "
        "does not include I_GLAccountLineItemRawData, so that payment is not in the sample.",
    )

    doc.add_heading("11. How the documents are correlated", level=1)
    add_table(
        doc,
        ["From", "To", "Keys"],
        [
            [row["from"], row["to"], "; ".join(row["keys"]) + ((". " + row["note"]) if row.get("note") else "")]
            for row in c["joinKeys"]
        ],
        font=8,
    )

    doc.add_heading("12. Extract calls for this reference", level=1)
    add_para(doc, "One ExtractCds call per entity. Example for the requisition:")
    add_para(
        doc,
        "GET /sap/opu/odata/sap/ZEVO_CDS_EXTRACT_SRV/ExtractCds"
        "?EntityName='C_PurchaseRequisitionItemDEX'"
        "&Filter='PurchaseRequisition eq ''0010001634'''"
        "&Format='jsonrows'&Skip='0'&Top='20'&$format=json",
        size=9,
    )
    add_table(
        doc,
        ["Entity", "Filter"],
        [
            ["C_PurchaseContractDEX", "PurchaseContract eq '4600000042'"],
            ["C_PurchaseContractItemDEX", "PurchaseContract eq '4600000042'"],
            ["C_PurchaseContractHistoryDEX", "PurchaseContract eq '4600000042'"],
            ["C_PurchaseOrderDEX", "PurchaseOrder eq '4500002148'"],
            ["C_PurchaseOrderItemDEX", "PurchaseOrder eq '4500002148'"],
            ["C_PurOrdAccountAssignmentDEX", "PurchaseOrder eq '4500002148'"],
            ["C_PurchaseOrderHistoryDEX", "PurchaseOrder eq '4500002148'"],
            ["I_GoodsMovementDocumentDEX", "PurchaseOrder eq '4500002148'"],
            ["C_SupplierInvoiceItemDEX", "PurchaseOrder eq '4500002148'"],
            ["C_SupplierInvoiceDEX", "SupplierInvoice eq '5100001601'"],
            ["I_BusinessPartner", "BusinessPartner eq '0001000579'"],
            ["I_GLAccount", "CompanyCode eq '1710' and GLAccount eq '0054400000'"],
        ],
        font=8,
    )

    doc.add_heading("13. Sample-data files", level=1)
    add_para(
        doc,
        "CSV delimiter is semicolon. Rebuild with python3 docs/samples/rtp/build_seed.py "
        "and python3 docs/samples/rtp/build_use_case_docx.py. Source extracts are in "
        "docs/samples/rtp/source. The rows for this chain are in docs/samples/rtp/seed.",
    )
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
