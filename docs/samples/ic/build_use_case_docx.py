#!/usr/bin/env python3
"""Build docs/INTERCOMPANY_RESALE.docx for the two live 1710 → 1010 resales.

These cases were posted in SAP. They are not a ZEVO extract pack.
Run from the repository root:

    python3 docs/samples/ic/build_use_case_docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOCS = ROOT.parents[1]
OUT = DOCS / "INTERCOMPANY_RESALE.docx"

NAVY = "1F4E79"
ALT = "F3F6F9"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
NAVY_RGB = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x55, 0x55, 0x55)
META = RGBColor(0x33, 0x33, 0x33)


def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, size=11, bold=False, space_after=8, space_before=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=BLACK)
    return paragraph


def shade_cell(cell, fill: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, size=8, bold=False, color=None, fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run("" if text is None else str(text))
    set_run_font(run, size=size, bold=bold, color=color or BLACK)
    if fill:
        shade_cell(cell, fill)
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def add_table(doc, headers, rows, font=8):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, size=font, bold=True, color=WHITE, fill=NAVY)
    repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for r_index, row in enumerate(rows):
        fill = ALT if r_index % 2 else "FFFFFF"
        for c_index, value in enumerate(row):
            set_cell_text(table.rows[r_index + 1].cells[c_index], value, size=font, fill=fill)
        prevent_row_split(table.rows[r_index + 1])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("1710 → 1010  ·  case 1 billing 90005786  ·  case 2 billing 90005787  ·  ")
    set_run_font(run, size=8, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    page_run = paragraph.add_run()
    set_run_font(page_run, size=8, color=GRAY)
    page_run._r.append(fld_begin)
    page_run._r.append(instr)
    page_run._r.append(fld_end)


def case_one(doc) -> None:
    doc.add_heading("Case 1. Resale with differences", level=1)
    add_para(
        doc,
        "Company 1710 sold 10 pieces of TG11 to company 1010. The purchase order "
        "and the customer invoice were posted, and the vendor invoice was paid. "
        "The two commercial prices were left different on purpose.",
    )

    doc.add_heading("1. Purpose", level=2)
    add_para(
        doc,
        "This case records an intercompany resale in which the documents do not "
        "agree. Germany bought in US dollars. The United States billed in euros. "
        "Tax was charged only on the German vendor invoice. The difference is the "
        "point of the case.",
    )

    doc.add_heading("2. Document flow", level=2)
    add_table(
        doc,
        ["Step", "Document", "What it carries"],
        [
            ["1", "Purchase order 4500002162", "Company 1010 buys 10 PC of TG11 from BP1710 at 27.00 USD"],
            ["2", "Goods receipt 5000002991", "10 PC received in plant 1010 on 13.10.2026, value 286.56 EUR"],
            ["3", "Billing 90005786", "Company 1710 invoices payer 10401010 for 203.76 EUR, tax 0.00"],
            ["4", "Vendor invoice 5100001607", "321.30 USD to vendor 0017401710. Reference 90005786"],
            ["5", "Outgoing payment in company 1010", "The vendor invoice was paid. The payment number was not recorded"],
        ],
    )
    add_para(
        doc,
        "The sales-order number and the outbound delivery number for this billing "
        "document were not captured. The billing document is the selling document "
        "that was kept.",
    )

    doc.add_heading("3. Business partner", level=2)
    add_table(
        doc,
        ["Role", "Number", "Name", "What it represents"],
        [
            ["Supplier on the purchase order", "BP1710", "Plant 1", "Plant 1710"],
            ["Vendor on the posted invoice", "0017401710", "SAP SE", "Company 1710"],
            ["Payer on the customer invoice", "10401010", "Buchungskreis 1010 als Geschäftspartner", "Company 1010"],
            ["Plant partner, not the payer", "10411010", "Werk 1010 als Geschäftspartner", "Plant 1010"],
        ],
    )
    add_para(
        doc,
        "The purchase order was entered with supplier BP1710. The invoice posted "
        "to vendor 0017401710. The customer invoice posted to payer 10401010. "
        "Customer 10411010 had no open item for this billing document.",
    )

    doc.add_heading("4. Source document", level=2)
    add_para(doc, "The source document on the buying side is purchase order 4500002162.")
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Purchase order", "4500002162"],
            ["Supplier entered", "BP1710"],
            ["Company code", "1010"],
            ["Purchasing organization", "1010"],
            ["Material", "TG11"],
            ["Quantity", "10 PC"],
            ["Net price", "27.00 USD per 1 PC"],
            ["Net amount", "270.00 USD"],
            ["Receiving plant", "1010, Plant 1 DE"],
            ["Document date", "13.10.2026"],
        ],
    )

    doc.add_heading("5. Follow-on document", level=2)
    add_para(
        doc,
        "The follow-on selling document is billing document 90005786. It invoices "
        "the same 10 pieces at a different price and in a different currency.",
    )
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Billing document", "90005786"],
            ["Billing date", "13.10.2026"],
            ["Payer", "10401010"],
            ["Material", "TG11"],
            ["Quantity", "10 PC"],
            ["Net value", "203.76 EUR"],
            ["Tax", "0.00"],
        ],
    )

    doc.add_heading("6. Goods movement", level=2)
    add_table(
        doc,
        ["Document", "Movement", "Quantity", "Value"],
        [
            ["5000002991", "101 goods receipt", "10 PC into plant 1010", "286.56 EUR"],
        ],
    )
    add_para(
        doc,
        "The goods issue in plant 1710 was posted before this receipt. Its material "
        "document number was not recorded.",
    )

    doc.add_heading("7. Invoice", level=2)
    add_table(
        doc,
        ["Invoice", "Company", "Partner", "Amount"],
        [
            ["5100001607", "1010", "Vendor 0017401710", "270.00 USD goods + 51.30 USD tax = 321.30 USD"],
            ["90005786", "1710", "Payer 10401010", "203.76 EUR, tax 0.00"],
        ],
    )
    add_para(
        doc,
        "Tax code V1, input tax at the full rate, added 19 percent on the vendor "
        "invoice. 51.30 is 19 percent of 270.00. The customer invoice has no tax. "
        "The vendor invoice reference is billing document 90005786.",
    )

    doc.add_heading("8. Accounting", level=2)
    add_para(
        doc,
        "The goods receipt in company 1010 valued the 10 pieces at 286.56 EUR. "
        "That is the local-currency value of the 270.00 USD purchase price. "
        "The vendor invoice then added 51.30 USD of tax. The customer invoice "
        "in company 1710 is an EUR document. Company 1710 translates it to USD "
        "at the rate on the billing date. The USD amount of that billing document "
        "was not recorded.",
    )

    doc.add_heading("9. Quantity and amount reconciliation", level=2)
    add_table(
        doc,
        ["Comparison", "Purchase side", "Sales side"],
        [
            ["Quantity", "10 PC", "10 PC"],
            ["Goods amount", "270.00 USD", "203.76 EUR"],
            ["Local value of the goods", "286.56 EUR", "203.76 EUR"],
            ["Tax", "51.30 USD", "0.00"],
            ["Amount paid or billed", "321.30 USD paid", "203.76 EUR billed"],
        ],
    )
    add_para(
        doc,
        "Quantity matches. Value does not. The euro value of the purchase, "
        "286.56 EUR, is 82.80 EUR higher than the customer invoice of 203.76 EUR.",
    )

    doc.add_heading("10. Clearing", level=2)
    add_para(
        doc,
        "The outgoing payment in company 1010 cleared vendor invoice 5100001607. "
        "The payment document number was not recorded. The incoming payment for "
        "billing document 90005786 was not posted in this session, so the "
        "receivable on customer 10401010 was left open.",
    )

    doc.add_heading("11. How the documents are correlated", level=2)
    add_table(
        doc,
        ["From", "To", "Key"],
        [
            ["Purchase order 4500002162", "Goods receipt 5000002991", "Purchase order and item on the material document"],
            ["Purchase order 4500002162", "Vendor invoice 5100001607", "Purchase-order reference, item 10, amount 270.00 USD"],
            ["Billing 90005786", "Vendor invoice 5100001607", "Invoice reference 90005786"],
            ["Billing 90005786", "Customer 10401010", "Payer on the billing document"],
        ],
    )

    doc.add_heading("12. Extract calls for this reference", level=2)
    add_para(
        doc,
        "There is no ZEVO extract for this case. The documents were posted in "
        "SAP and read from the transactions. An extract taken later would filter "
        "purchase order 4500002162, billing document 90005786, material document "
        "5000002991, and supplier invoice 5100001607.",
    )

    doc.add_heading("13. Sample-data files", level=2)
    add_para(
        doc,
        "There is no sample-data pack. The numbers in this case are the numbers "
        "posted in the system.",
    )


def case_two(doc) -> None:
    doc.add_heading("Case 2. Purchase, then resale at a higher price", level=1)
    add_para(
        doc,
        "Company 1710 bought 5 pieces of TG11 from a domestic supplier, then sold "
        "the same pieces to company 1010 at 25.00 EUR each. The goods amount is "
        "125.00 EUR on both the customer invoice and the vendor invoice.",
    )

    doc.add_heading("1. Purpose", level=2)
    add_para(
        doc,
        "This case shows a completed buy and a completed resale. The US company "
        "earns a markup: it bought at 13.50 USD per piece and sold at 25.00 EUR "
        "per piece. The intercompany goods price is the same on both sides. "
        "German input tax remains only on the vendor invoice.",
    )

    doc.add_heading("2. Document flow", level=2)
    add_table(
        doc,
        ["Step", "Document", "What it carries"],
        [
            ["1", "Contract 4600000041 item 30", "Outline agreement for TG11 at 13.50 USD, supplier 1000559"],
            ["2", "Purchase order 4500002164", "Release of 5 PC into plant 1710"],
            ["3", "Goods receipt 5000003010", "5 PC received in plant 1710, storage location 171A"],
            ["4", "Supplier invoice 5100001609", "67.50 USD. Accounting document 5100000005, type RE"],
            ["5", "Payment 1500000004", "67.50 USD to vendor 1000559 on 19.10.2026"],
            ["6", "Sales order 6323", "5 PC at 25.00 EUR to customer 10401010. Value 125.00 EUR"],
            ["7", "Purchase order 4500002165", "Company 1010 buys the same 5 PC at 25.00 EUR"],
            ["8", "Delivery 80006442", "Goods issue 4900008971 from plant 1710"],
            ["9", "Goods receipt 5000003011", "5 PC received in plant 1010"],
            ["10", "Billing 90005787", "125.00 EUR to payer 10401010"],
            ["11", "Vendor invoice 5100001610", "148.75 EUR. Reference INV.90005787"],
            ["12", "Payment 1500000001", "148.75 EUR paid in company 1010"],
            ["13", "Receipt 1400000002", "125.00 EUR received in company 1710"],
        ],
    )

    doc.add_heading("3. Business partner", level=2)
    add_table(
        doc,
        ["Role", "Number", "Name", "Where it is used"],
        [
            ["Domestic supplier", "1000559", "EVOLVER DOMESTIC SUPPLIER 1", "Purchase order 4500002164 and payment 1500000004"],
            ["Supplying plant as supplier", "BP1710", "Plant 1", "Purchase order 4500002165"],
            ["Company 1710 as vendor", "0017401710", "SAP SE", "Vendor that carries company 1710. Used on the earlier invoice 5100001607"],
            ["Company 1010 as customer", "10401010", "Buchungskreis 1010 als Geschäftspartner", "Sold-to and payer on sales order 6323 and billing 90005787"],
        ],
    )
    add_para(
        doc,
        "Sales organization 1710, distribution channel 10, division 00. "
        "Customer 10401010 is a German partner, so the sales order currency is EUR. "
        "Company 1710 still posts its books in USD.",
    )

    doc.add_heading("4. Source document", level=2)
    add_para(doc, "The source document is the release purchase order against the outline agreement.")
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Purchase order", "4500002164"],
            ["Contract", "4600000041 item 30"],
            ["Vendor", "1000559"],
            ["Company code", "1710"],
            ["Purchasing organization", "1710"],
            ["Purchasing group", "001"],
            ["Material", "TG11, Trad.Good 11, PD, Reg.Trading"],
            ["Quantity", "5 PC"],
            ["Net price", "13.50 USD"],
            ["Plant and storage location", "1710 / 171A"],
        ],
    )

    doc.add_heading("5. Follow-on document", level=2)
    add_para(doc, "The resale has two follow-on commercial documents, one in each company.")
    add_table(
        doc,
        ["Document", "Company", "Partner", "Price"],
        [
            ["Sales order 6323", "1710", "Customer 10401010", "5 × 25.00 EUR = 125.00 EUR"],
            ["Purchase order 4500002165", "1010", "Supplier BP1710", "5 × 25.00 EUR = 125.00 EUR"],
        ],
    )
    add_para(
        doc,
        "Purchase order 4500002165 uses purchasing organization 1010, purchasing "
        "group 001, company code 1010, currency EUR, and plant 1010. The header "
        "exchange rate is 1.00000 because the document currency and the company "
        "currency are both EUR. The order reference is sales order 6323. "
        "The sales order was released in credit control area 1000 before the delivery.",
    )

    doc.add_heading("6. Goods movement", level=2)
    add_table(
        doc,
        ["Document", "What moved", "Quantity", "Plant"],
        [
            ["5000003010", "Goods receipt against purchase order 4500002164, movement 101", "5 PC", "1710 / 171A"],
            ["4900008971", "Goods issue for delivery 80006442", "5 PC", "1710"],
            ["5000003011", "Goods receipt against purchase order 4500002165, movement 101", "5 PC", "1010 / 101A"],
        ],
    )

    doc.add_heading("7. Invoice", level=2)
    add_table(
        doc,
        ["Invoice", "Company", "Partner", "Amount"],
        [
            ["5100001609", "1710", "Vendor 1000559", "67.50 USD"],
            ["90005787", "1710", "Payer 10401010", "125.00 EUR"],
            ["5100001610", "1010", "Vendor of purchase order 4500002165", "125.00 EUR goods + 23.75 EUR tax = 148.75 EUR"],
        ],
    )
    add_para(
        doc,
        "Vendor invoice 5100001610 is dated 20.10.2026. Its reference is "
        "INV.90005787. Calculate tax was selected. Tax code V1 added 23.75 EUR, "
        "which is 19 percent of 125.00 EUR. The customer invoice is 125.00 EUR "
        "and has no tax.",
    )

    doc.add_heading("8. Accounting", level=2)
    add_para(doc, "Payment 1500000004 in company 1710, fiscal year 2026, period 10, currency USD.")
    add_table(
        doc,
        ["Line", "Posting key", "Account", "Amount"],
        [
            ["1", "50", "11002000 Bank2 Main Account", "67.50- USD"],
            ["2", "25", "1000559 EVOLVER DOMESTIC SUPPLIER 1", "67.50 USD"],
        ],
    )
    add_para(
        doc,
        "Document date and posting date are 19.10.2026. The supplier invoice "
        "5100001609 posted to accounting document 5100000005, type RE, for the "
        "same 67.50 USD. Payment 1500000001 cleared the 148.75 EUR vendor invoice "
        "in company 1010. Receipt 1400000002 cleared the 125.00 EUR customer "
        "invoice in company 1710.",
    )

    doc.add_heading("9. Quantity and amount reconciliation", level=2)
    add_table(
        doc,
        ["Comparison", "US purchase", "Intercompany sale", "German purchase"],
        [
            ["Quantity", "5 PC", "5 PC", "5 PC"],
            ["Price per piece", "13.50 USD", "25.00 EUR", "25.00 EUR"],
            ["Goods amount", "67.50 USD", "125.00 EUR", "125.00 EUR"],
            ["Tax", "None on the paid invoice", "None", "23.75 EUR"],
            ["Amount cleared", "67.50 USD", "125.00 EUR", "148.75 EUR"],
        ],
    )
    add_para(
        doc,
        "The intercompany goods amount matches at 125.00 EUR. The amount Germany "
        "paid is higher by the 23.75 EUR of input tax. The markup is the difference "
        "between the sale price of 25.00 EUR and the purchase cost of 13.50 USD, "
        "after company 1710 translates the euro invoice into dollars.",
    )

    doc.add_heading("10. Clearing", level=2)
    add_table(
        doc,
        ["Payment", "Company", "Clears", "Amount"],
        [
            ["1500000004", "1710", "Supplier invoice 5100001609, vendor 1000559", "67.50 USD"],
            ["1500000001", "1010", "Vendor invoice 5100001610", "148.75 EUR"],
            ["1400000002", "1710", "Billing 90005787, customer 10401010", "125.00 EUR"],
        ],
    )
    add_para(
        doc,
        "Each open item was selected in full and posted at a balance of 0.00. "
        "A 3 percent cash discount of 2.03 USD was proposed on invoice 5100001609 "
        "and was removed, so the vendor was paid the full 67.50 USD.",
    )

    doc.add_heading("11. How the documents are correlated", level=2)
    add_table(
        doc,
        ["From", "To", "Key"],
        [
            ["Contract 4600000041 item 30", "Purchase order 4500002164", "Outline agreement referenced on the purchase-order item"],
            ["Purchase order 4500002164", "Goods receipt 5000003010 and invoice 5100001609", "Purchase order number"],
            ["Invoice 5100001609", "Payment 1500000004", "Vendor open item 5100000005, amount 67.50 USD"],
            ["Sales order 6323", "Delivery 80006442 and billing 90005787", "Sales order, then the delivery number on the billing document"],
            ["Sales order 6323", "Purchase order 4500002165", "Reference 6323 on the purchase order"],
            ["Purchase order 4500002165", "Goods receipt 5000003011 and invoice 5100001610", "Purchase order number"],
            ["Billing 90005787", "Vendor invoice 5100001610", "Reference INV.90005787"],
            ["Billing 90005787", "Receipt 1400000002", "Customer 10401010 and amount 125.00 EUR"],
        ],
    )

    doc.add_heading("12. Extract calls for this reference", level=2)
    add_para(
        doc,
        "There is no ZEVO extract for this case. The documents were posted in "
        "SAP during the session. An extract taken later would start from purchase "
        "order 4500002164 and sales order 6323, then follow purchase order "
        "4500002165, delivery 80006442, billing document 90005787, and supplier "
        "invoice 5100001610.",
    )

    doc.add_heading("13. Sample-data files", level=2)
    add_para(
        doc,
        "There is no sample-data pack. Rebuild this document with "
        "python3 docs/samples/ic/build_use_case_docx.py. The output is "
        "docs/INTERCOMPANY_RESALE.docx.",
    )


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = NAVY_RGB
        style.font.size = Pt(size)
        style.font.bold = True

    core = doc.core_properties
    core.title = "Intercompany resale — company 1710 to company 1010"
    core.subject = "Two live resale scenarios for TG11"
    core.category = "Intercompany resale"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Intercompany resale")
    set_run_font(run, size=22, bold=True, color=NAVY_RGB)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    run = subtitle.add_run("Two end-to-end scenarios")
    set_run_font(run, size=14, color=NAVY_RGB)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run(
        "Company 1710 (United States, USD) sells trading good TG11 to company 1010 (Germany, EUR)\n"
        "Case 1 keeps the price, currency, and tax differences\n"
        "Case 2 buys from a domestic supplier, then resells at a higher price\n"
        "Posted in SAP. Not taken from a ZEVO extract"
    )
    set_run_font(run, size=10, color=META)

    add_para(
        doc,
        "Both cases use the same product and the same two companies. Case 1 is "
        "the resale in which the documents were left inconsistent. Case 2 is the "
        "resale in which company 1710 first buys the goods and then sells them "
        "to company 1010 at one agreed euro price.",
    )
    add_table(
        doc,
        ["", "Case 1", "Case 2"],
        [
            ["Quantity", "10 PC", "5 PC"],
            ["US purchase", "Not part of this posting", "13.50 USD per piece, 67.50 USD"],
            ["Price charged to Germany", "27.00 USD on the purchase order, 203.76 EUR on the invoice", "25.00 EUR on both sides, 125.00 EUR"],
            ["Tax on the German invoice", "51.30 USD", "23.75 EUR"],
            ["Result", "Quantity matches. Value does not", "Goods value matches. Tax remains only in Germany"],
        ],
    )

    case_one(doc)
    case_two(doc)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
