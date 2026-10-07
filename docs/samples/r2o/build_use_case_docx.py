#!/usr/bin/env python3
"""Build docs/REQUISITION_TO_ORDER.docx from correlation.json.

    python3 docs/samples/r2o/build_use_case_docx.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOCS = ROOT.parents[1]
OUT = DOCS / "REQUISITION_TO_ORDER.docx"

sys.path.insert(0, str(ROOT.parent / "c2p"))
from build_use_case_docx import BLACK, add_para, add_table, set_run_font  # noqa: E402


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Requisition 10001624  ·  PO 4500002147  ·  ")
    set_run_font(run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    page_run = paragraph.add_run()
    set_run_font(page_run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    page_run._r.append(begin)
    page_run._r.append(instr)
    page_run._r.append(end)


def main() -> None:
    c = json.loads((ROOT / "correlation.json").read_text(encoding="utf-8"))
    po = c["purchaseOrder"]
    invoice = c["supplierInvoice"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        style.font.size = Pt(size)
        style.font.bold = True

    core = doc.core_properties
    core.title = "Requisition to Order — 10001624 / 4500002147"
    core.subject = "ZEVO CDS extract use case with sample data"
    core.category = "Requisition to Order"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_run_font(title.add_run("Requisition to Order"), size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    set_run_font(subtitle.add_run("Use case with sample data"), size=14, color=RGBColor(0x1F, 0x4E, 0x79))
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    set_run_font(
        meta.add_run(
            "Reference documents: purchase requisition 10001624 and purchase order 4500002147\n"
            "SAP S/4HANA DEX extracts · ZEVO_CDS_EXPLORER_2_FILE · run 20261007_121746\n"
            "Company code 1710 · purchasing organization 1710 · USD"
        ),
        size=10,
        color=RGBColor(0x33, 0x33, 0x33),
    )

    doc.add_heading("1. Purpose", level=1)
    add_para(
        doc,
        "This use case shows how a released purchase requisition becomes a purchase order. "
        "The sample rows are the ZEVO extracts for requisition 10001624 and purchase order "
        "4500002147. Both items are account-assigned to the same fixed asset.",
    )
    add_para(
        doc,
        "ZEVO writes one CDS entity per file. The requisition item in this extract does not "
        "store the purchase-order number. The purchase-order item stores the requisition "
        "number and item. Follow that direction. Dates are YYYYMMDD. Document numbers keep "
        "their leading zeros.",
    )

    doc.add_heading("2. Document flow", level=1)
    add_table(
        doc,
        ["Step", "Document", "What it carries"],
        [
            ["1", "Requisition 0010001624", "Two released items, supplier 0001000579, account assignment A"],
            ["2", "Purchase order 4500002147", "Same two items, same quantities and prices, total 43,000.00 USD"],
            ["3", "Account assignment", "Both items to asset 000000600004 and G/L 0016014000"],
            ["4", "Goods receipts 5000002950 and 5000002951", "10 of 20 received on each item"],
            ["5", "Supplier invoice 5100001600", "10 of 20 invoiced, gross 21,500.00 USD"],
        ],
    )
    add_para(
        doc,
        "Steps 4 and 5 are the follow-on documents already posted against the order. "
        "Half of each item is still open. This extract set has no journal file, so the "
        "payment of the invoice is not in the sample.",
    )

    doc.add_heading("3. Why this requisition", level=1)
    add_para(
        doc,
        "Requisition 0010001624 is the only requisition in the extract whose items appear "
        "on a purchase order. Requisition 0010001573 item 00010 (material RM2_DS, quantity "
        "150, price 5.25, processing status N) has no purchase order in this run.",
    )

    doc.add_heading("4. Supplier", level=1)
    add_para(
        doc,
        f"Supplier {c['supplier']['supplier']} is {c['supplier']['name']}. "
        "The same number is the business partner, the supplier on the requisition, "
        "the supplier on the purchase order, and the invoicing party.",
    )

    doc.add_heading("5. Purchase requisition 0010001624", level=1)
    first = c["requisitionItems"][0]
    add_para(
        doc,
        f"Type {first['requisitionType']}, company code {first['companyCode']}, "
        f"plant {first['plant']}, purchasing group {first['purchasingGroup']}. "
        f"Created by {first['createdByUser']} on {first['creationDate']}. "
        f"Release status {first['releaseStatus']}, release date {first['releaseDate']}. "
        "Processing status B means a purchase order has been created.",
    )
    add_para(doc, "C_PurchaseRequisitionItemDEX", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Item", "Text", "Qty", "Unit", "Price", "Net amount", "Delivery date", "Supplier", "Acct asg.", "Release", "Status"],
        [
            [
                item["purchaseRequisitionItem"],
                item["text"],
                item["requestedQuantity"],
                item["unit"],
                item["price"],
                item["netAmount"],
                item["deliveryDate"],
                item["supplier"],
                item["accountAssignmentCategory"],
                item["releaseStatus"],
                item["processingStatus"],
            ]
            for item in c["requisitionItems"]
        ],
        font=8,
    )
    add_para(
        doc,
        "40,000.00 + 3,000.00 = 43,000.00. Neither item has a material number. "
        "The short text is the description. Account assignment category A is asset.",
    )

    doc.add_heading("6. Purchase order 4500002147", level=1)
    add_para(
        doc,
        f"Type {po['purchaseOrderType']}, date {po['purchaseOrderDate']}, created by "
        f"{po['createdByUser']}. Payment terms {po['paymentTerms']}. Header total "
        f"{po['totalAmount']} {po['documentCurrency']}. Processing status "
        f"{po['purchasingProcessingStatus']}.",
    )
    add_para(doc, "C_PurchaseOrderItemDEX — the link back to the requisition", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["PO item", "Requisition", "Req. item", "Text", "Qty", "Net price", "Net amount", "Acct asg.", "Plant"],
        [
            [
                item["purchaseOrderItem"],
                item["purchaseRequisition"],
                item["purchaseRequisitionItem"],
                item["text"],
                item["orderQuantity"],
                item["netPrice"],
                item["netAmount"],
                item["accountAssignmentCategory"],
                item["plant"],
            ]
            for item in po["items"]
        ],
        font=8,
    )
    add_para(
        doc,
        "Item 00010 of the order is item 00010 of the requisition. Item 00020 matches "
        "item 00020. Quantity, price, plant, supplier, and account assignment are copied "
        "unchanged. Goods receipt and invoice are expected on both items.",
    )

    doc.add_heading("7. Asset account assignment", level=1)
    add_para(
        doc,
        f"C_PurOrdAccountAssignmentDEX assigns both items to master fixed asset {c['asset']}, "
        f"assignment number 01, profit center YB110. The G/L account is "
        f"{c['glAccount']['glAccount']} (external {c['glAccount']['external']}), "
        f"group {c['glAccount']['group']}, reconciliation type "
        f"{c['glAccount']['reconciliationAccountType']}. That is the asset reconciliation account.",
    )
    add_table(
        doc,
        ["PO item", "Assignment", "Asset", "G/L account", "Quantity", "Profit center"],
        [
            [
                row["purchaseOrderItem"],
                row["accountAssignmentNumber"],
                row["masterFixedAsset"],
                row["glAccount"],
                row["quantity"],
                row["profitCenter"],
            ]
            for row in c["accountAssignments"]
        ],
    )

    doc.add_heading("8. Follow-on documents already on the order", level=1)
    add_para(
        doc,
        "C_PurchaseOrderHistoryDEX type 1, category E, is the goods receipt. The history "
        "document is the material document. Type 2, category Q, is the invoice receipt. "
        "The history document is the supplier invoice.",
    )
    add_para(doc, "Goods receipts", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Material document", "Year", "PO item", "Qty", "Movement", "Posting date"],
        [
            [
                row["materialDocument"],
                row["materialDocumentYear"],
                row["purchaseOrderItem"],
                row["quantity"],
                row["goodsMovementType"],
                row["postingDate"],
            ]
            for row in c["goodsReceipts"]
        ],
    )
    add_para(doc, "Supplier invoice", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Invoice", "Party reference", "Item", "PO item", "Qty", "Amount", "Posting date", "Status"],
        [
            [
                invoice["supplierInvoice"],
                invoice["reference"],
                item["supplierInvoiceItem"],
                item["purchaseOrderItem"],
                item["quantity"],
                item["amount"],
                invoice["postingDate"],
                invoice["status"],
            ]
            for item in invoice["items"]
        ],
    )
    add_para(
        doc,
        f"Invoice {invoice['supplierInvoice']} gross {invoice['grossAmount']} "
        f"{invoice['currency']} equals 20,000.00 + 1,500.00. Status {invoice['status']} "
        "with a posted invoice means the invoice is posted. Each line invoices 10 of the "
        "20 ordered.",
    )

    doc.add_heading("9. Quantity reconciliation", level=1)
    add_table(
        doc,
        ["Req. item", "PO item", "Text", "Ordered", "Order amount", "Received", "Invoiced qty", "Invoiced amount", "Still open qty", "Still open amount"],
        [
            [
                row["purchaseRequisitionItem"],
                row["purchaseOrderItem"],
                row["text"],
                row["orderedQuantity"],
                row["orderedAmount"],
                row["receivedQuantity"],
                row["invoicedQuantity"],
                row["invoicedAmount"],
                row["openQuantity"],
                row["openAmount"],
            ]
            for row in c["quantityReconciliation"]
        ],
        font=7,
    )
    add_para(
        doc,
        "Received quantity equals invoiced quantity. The open amount is the quantity "
        "not yet received: 10 laptops (20,000.00) and 10 keyboards (1,500.00).",
    )

    doc.add_heading("10. How the documents are correlated", level=1)
    add_table(
        doc,
        ["From", "To", "Keys"],
        [
            [row["from"], row["to"], "; ".join(row["keys"]) + ((". " + row["note"]) if row.get("note") else "")]
            for row in c["joinKeys"]
        ],
        font=8,
    )

    doc.add_heading("11. Extract calls for this reference", level=1)
    add_para(
        doc,
        "One ExtractCds call per entity. Example for the requisition items:",
    )
    add_para(
        doc,
        "GET /sap/opu/odata/sap/ZEVO_CDS_EXTRACT_SRV/ExtractCds"
        "?EntityName='C_PurchaseRequisitionItemDEX'"
        "&Filter='PurchaseRequisition eq ''0010001624'''"
        "&Format='jsonrows'&Skip='0'&Top='20'&$format=json",
        size=9,
    )
    add_table(
        doc,
        ["Entity", "Filter"],
        [
            ["C_PurchaseOrderItemDEX", "PurchaseRequisition eq '0010001624'"],
            ["C_PurchaseOrderDEX", "PurchaseOrder eq '4500002147'"],
            ["C_PurOrdAccountAssignmentDEX", "PurchaseOrder eq '4500002147'"],
            ["C_PurchaseOrderHistoryDEX", "PurchaseOrder eq '4500002147'"],
            ["I_GoodsMovementDocumentDEX", "PurchaseOrder eq '4500002147'"],
            ["C_SupplierInvoiceItemDEX", "PurchaseOrder eq '4500002147'"],
            ["C_SupplierInvoiceDEX", "SupplierInvoice eq '5100001600'"],
            ["I_BusinessPartner", "BusinessPartner eq '0001000579'"],
            ["I_GLAccount", "CompanyCode eq '1710' and GLAccount eq '0016014000'"],
        ],
        font=8,
    )

    doc.add_heading("12. Sample-data files", level=1)
    add_para(
        doc,
        "CSV delimiter is semicolon. Rebuild with python3 docs/samples/r2o/build_seed.py "
        "and python3 docs/samples/r2o/build_use_case_docx.py. Source extracts are in "
        "docs/samples/r2o/source. The rows for this requisition and order are in "
        "docs/samples/r2o/seed.",
    )

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
