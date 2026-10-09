#!/usr/bin/env python3
"""Build docs/REQUISITION_TO_PAYMENT_ER.pptx.

Same layout as the contract-to-payment diagram, with the requisition in front.
Sample is requisition 0010001634, contract 4600000042, PO 4500002148.

    python3 docs/samples/c2p/build_requisition_er_pptx.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_er_pptx import AMBER, INK, LINE, MUTED, NAVY, PALE, PLUM, TEAL, WHITE
from build_er_pptx import arrow, box, footer, label, set_run

OUT = Path(__file__).resolve().parents[2] / "REQUISITION_TO_PAYMENT_ER.pptx"
RUST = RGBColor(0x8C, 0x3A, 0x2F)


def main() -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = PALE
    background.line.fill.background()

    label(slide, Inches(0.35), Inches(0.18), Inches(12), Inches(0.36), "Requisition to Payment", size=26, color=NAVY)
    label(
        slide,
        Inches(0.35),
        Inches(0.54),
        Inches(12.5),
        Inches(0.26),
        "Entity diagram  ·  requisition 10001634  ·  contract 4600000042  ·  PO 4500002148  ·  payment 1500000002",
        size=13,
        bold=False,
        color=MUTED,
    )

    columns = [
        (RUST, "Requisition", [
            ("RequisitionItem", ["PK Requisition + Item", "Status K = contract"]),
            ("BusinessPartner", ["Supplier = BusinessPartner", "Also the invoicing party"]),
        ]),
        (NAVY, "Contract", [
            ("PurchaseContract", ["PK PurchaseContract", "Supplier, CompanyCode"]),
            ("PurchaseContractItem", ["PK Contract + Item", "Target qty and price"]),
            ("ContractHistory", ["ReleaseOrder = PO", "ReleaseOrderItem"]),
        ]),
        (TEAL, "Purchase order", [
            ("PurchaseOrder", ["PK PurchaseOrder", "Supplier, CompanyCode"]),
            ("PurchaseOrderItem", ["FK PurchaseContract + Item", "No requisition number here"]),
            ("PurchaseOrderHistory", ["Type 1 = goods receipt", "Type 2 = invoice"]),
        ]),
        (AMBER, "Receipt and invoice", [
            ("GoodsMovement", ["PK MaterialDocument", "Year + Item, movement 101"]),
            ("SupplierInvoice", ["PK Invoice + Year", "+ CompanyCode"]),
            ("SupplierInvoiceItem", ["FK PO + Item", "Quantity and amount"]),
        ]),
        (PLUM, "Journal and payment", [
            ("JournalLine WE", ["ReferenceDocumentType MKPF", "ReferenceDocument = mat. doc"]),
            ("JournalLine RE", ["ReferenceDocumentType RMRP", "ReferenceDocument = invoice"]),
            ("JournalLine KZ", ["Payment document", "ClearingAccountingDocument"]),
        ]),
    ]

    lefts = [0.28, 2.9, 5.52, 8.14, 10.76]
    width = 2.4
    for index, (color, heading, cards) in enumerate(columns):
        left = Inches(lefts[index])
        label(slide, left, Inches(0.92), Inches(width), Inches(0.26), heading, size=13, color=color)
        top = 1.24
        height = 1.7 if len(cards) == 2 else 1.55
        gap = 0.1 if len(cards) == 2 else 0.08
        if len(cards) == 2:
            top = 1.7
            height = 1.85
        for title, body in cards:
            box(slide, left, Inches(top), Inches(width), Inches(height), color, title, body)
            top += height + gap
        if index < len(columns) - 1:
            arrow(slide, Inches(lefts[index] + width + 0.02), Inches(3.4), Inches(0.18), Inches(0.16))

    footer(slide, "The requisition item does not store the contract number. The PO item and the contract history do.")

    slide = presentation.slides.add_slide(blank)
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = WHITE
    background.line.fill.background()
    label(slide, Inches(0.4), Inches(0.2), Inches(12), Inches(0.34), "How the entities connect", size=26, color=NAVY)
    label(
        slide,
        Inches(0.4),
        Inches(0.54),
        Inches(12.4),
        Inches(0.26),
        "With a contract: requisition 10001634, contract 4600000042, PO 4500002148. A direct requisition skips the contract.",
        size=13,
        bold=False,
        color=MUTED,
    )

    rows = [
        ["Requisition → contract", "Same text, supplier, plant, quantity. Status K", "0010001634 → 4600000042"],
        ["Contract item → history", "PurchaseContract + PurchaseContractItem", "ReleaseOrder 4500002148"],
        ["Contract item → PO item", "PurchaseContract + PurchaseContractItem", "4600000042 / 00010 and 00020"],
        ["History type 1 → goods receipt", "PurchasingHistoryDocument = MaterialDocument", "5000002952"],
        ["History type 2 → invoice", "PurchasingHistoryDocument = SupplierInvoice", "5100001601"],
        ["Material document → journal", "ReferenceDocumentType MKPF", "FI 5000000004"],
        ["Invoice → journal", "ReferenceDocumentType RMRP", "FI 5100000003"],
        ["Invoice journal → payment", "ClearingAccountingDocument", "KZ 1500000002, 221.00 USD"],
        ["Direct requisition → PO", "PurchaseRequisition + Item on the PO item", "0010001624 → 4500002147"],
        ["That invoice → payment", "ClearingAccountingDocument", "KZ 1500000001, 21,500.00 USD"],
    ]
    table_shape = slide.shapes.add_table(1 + len(rows), 3, Inches(0.4), Inches(0.98), Inches(12.5), Inches(5.7))
    table = table_shape.table
    table.columns[0].width = Inches(3.6)
    table.columns[1].width = Inches(4.8)
    table.columns[2].width = Inches(4.1)
    for col, text in enumerate(["Relationship", "Key", "In this sample"]):
        cell = table.cell(0, col)
        cell.text = ""
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_run(cell.text_frame.paragraphs[0].add_run(), text, 12, True, WHITE)
    for r_index, row in enumerate(rows):
        for col, text in enumerate(row):
            cell = table.cell(r_index + 1, col)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALE if r_index % 2 == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            set_run(cell.text_frame.paragraphs[0].add_run(), text, 12, False, INK)

    footer(slide, "Status K means the requisition was converted to a contract. Status B means it was converted straight to a purchase order.")
    presentation.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
