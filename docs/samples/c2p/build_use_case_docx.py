#!/usr/bin/env python3
"""Build docs/CONTRACT_TO_PAYMENT.docx from correlation.json and the seed CSVs.

Requires python-docx. Run from the repository root:

    python3 docs/samples/c2p/build_use_case_docx.py
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOCS = ROOT.parents[1]
SEED = ROOT / "seed"
OUT = DOCS / "CONTRACT_TO_PAYMENT.docx"

NAVY = "1F4E79"
ALT = "F3F6F9"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

EMPTY = {
    "",
    "0",
    "0.00",
    "0.000",
    "0.00000",
    "0000",
    "00000",
    "000000",
    "00000000",
    "0000000000",
    "00000000000000000000000000000000",
}


def load_csv(name: str) -> list[dict[str, str]]:
    path = SEED / name
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


def filled(value: str) -> bool:
    return (value or "").strip() not in EMPTY


def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, size=11, bold=False, space_after=8, space_before=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=BLACK)
    return paragraph


def shade_cell(cell, fill: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, size=8, bold=False, color=None, fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run("" if text is None else str(text))
    set_run_font(run, size=size, bold=bold, color=color or BLACK)
    if fill:
        shade_cell(cell, fill)
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def repeat_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


def add_table(doc, headers: list[str], rows: list[list[str]], font=8):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, size=font, bold=True, color=WHITE, fill=NAVY)
    repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for r_index, row in enumerate(rows):
        fill = ALT if r_index % 2 else "FFFFFF"
        for c_index, value in enumerate(row):
            set_cell_text(table.rows[r_index + 1].cells[c_index], value, size=font, fill=fill)
        prevent_row_split(table.rows[r_index + 1])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_field_table(doc, pairs: list[tuple[str, str]]):
    add_table(doc, ["Field", "Value"], [[name, value] for name, value in pairs], font=9)


def nonempty_pairs(row: dict[str, str], fields: list[str] | None = None) -> list[tuple[str, str]]:
    keys = fields or list(row.keys())
    return [(key, row[key]) for key in keys if key in row and filled(row[key])]


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Contract 4600000041  ·  PO 4500002146  ·  extract 20261006_094549  ·  ")
    set_run_font(run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    page_run = paragraph.add_run()
    set_run_font(page_run, size=8, color=RGBColor(0x55, 0x55, 0x55))
    page_run._r.append(fld_begin)
    page_run._r.append(instr)
    page_run._r.append(fld_end)


def main() -> None:
    correlation = json.loads((ROOT / "correlation.json").read_text(encoding="utf-8"))
    contract = correlation["contract"]
    po = correlation["purchaseOrder"]
    payment = correlation["payment"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        style.font.size = Pt(size)
        style.font.bold = True

    core = doc.core_properties
    core.title = "Contract to Payment — 4600000041 / 4500002146"
    core.subject = "ZEVO CDS extract use case with sample data"
    core.category = "Contract to Payment"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Contract to Payment")
    set_run_font(run, size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    run = subtitle.add_run("Use case with sample data")
    set_run_font(run, size=14, color=RGBColor(0x1F, 0x4E, 0x79))

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run(
        "Reference documents: purchase contract 4600000041 and purchase order 4500002146\n"
        "SAP S/4HANA DEX extracts · ZEVO_CDS_EXPLORER_2_FILE · run 20261006_094549\n"
        "Company code 1710 · purchasing organization 1710 · USD"
    )
    set_run_font(run, size=10, color=RGBColor(0x33, 0x33, 0x33))

    doc.add_heading("1. Purpose", level=1)
    add_para(
        doc,
        "This use case shows how documents from a quantity contract are correlated "
        "through the release purchase order, goods receipt, supplier invoice, and the "
        "universal journal. The sample rows below are taken from the ZEVO file extracts "
        "for contract 4600000041 and purchase order 4500002146.",
    )
    add_para(
        doc,
        "ZEVO writes one CDS entity per file and one entity per ExtractCds call. It does "
        "not join documents. The client correlates the files with the keys in section 12. "
        "Full extracts are in docs/samples/c2p/source. The rows used here are in "
        "docs/samples/c2p/seed. Dates in the tables are YYYYMMDD, as stored in the CSV.",
    )

    doc.add_heading("2. Document flow", level=1)
    add_para(doc, "Reference chain for this use case:")
    add_table(
        doc,
        ["Step", "Document", "What it carries for this case"],
        [
            ["1", "Contract 4600000041", "Quantity contract MK, supplier 0001000559, three items"],
            ["2", "PO 4500002146", "Release of contract items 00020 (TG20) and 00030 (TG11)"],
            ["3", "Material documents 5000002931, 5000002940, 5000002941", "Goods receipts, movement type 101"],
            ["4", "Supplier invoices 5100001598 and 5100001599", "Invoice of the received quantities, 160.50 USD"],
            ["5", "Accounting documents 5000000001–5000000003 and 5100000000–5100000001", "Stock, GR/IR, and vendor items on ledger 0L"],
            ["6", "Payment", "Not in this extract. Vendor account 21100000 remains open for 160.50 USD"],
        ],
    )
    add_para(
        doc,
        "Contract item 00010 (TG10, target quantity 100) has no release order. "
        "There is no inbound delivery: DeliveryDocumentItem is 000000 on every history row. "
        "Both PO items have a blank account-assignment category, and "
        "C_PurOrdAccountAssignmentDEX contains a header and zero rows.",
    )

    doc.add_heading("3. Why these two documents", level=1)
    add_para(doc, correlation["whyThisSeed"])
    add_para(
        doc,
        "Contracts 4600000031 and 4600000040 are in the same contract extract and have no "
        "release rows. Purchase orders 4500002142 through 4500002145 release contract "
        "4600000038, which is not in the contract header or item files. Only 4500002143 "
        "has a goods receipt, and none of those orders has a supplier invoice in this extract.",
    )

    doc.add_heading("4. Supplier", level=1)
    add_para(
        doc,
        "Supplier 0001000559 is EVOLVER DOMESTIC SUPPLIER 1. The same number is the "
        "business partner and the supplier. Account group SUPL. Sample fields:",
    )
    bp = load_csv("I_BUSINESSPARTNER.csv")[0]
    sup = load_csv("I_BUSINESSPARTNERSUPPLIERDEX.csv")[0]
    add_para(doc, "I_BusinessPartner", size=11, bold=True, space_after=4)
    add_field_table(
        doc,
        nonempty_pairs(
            bp,
            [
                "BUSINESSPARTNER",
                "BUSINESSPARTNERCATEGORY",
                "BUSINESSPARTNERFULLNAME",
                "ORGANIZATIONBPNAME1",
                "BUSINESSPARTNERGROUPING",
                "SEARCHTERM1",
            ],
        ),
    )
    add_para(doc, "I_BusinessPartnerSupplierDEX", size=11, bold=True, space_after=4)
    add_field_table(doc, nonempty_pairs(sup))

    doc.add_heading("5. Purchase contract 4600000041", level=1)
    add_para(
        doc,
        "Quantity contract (type MK, purchasing-document category K), created by ARIA "
        "on 20261002. Validity 20261002 to 20271231. Payment terms 0003. Currency USD. "
        "Goods receipt and invoice are expected on every item.",
    )
    add_para(doc, "C_PurchaseContractDEX — header", size=11, bold=True, space_after=4)
    header_row = load_csv("C_PURCHASECONTRACTDEX.csv")[0]
    add_field_table(doc, nonempty_pairs(header_row))

    add_para(doc, "C_PurchaseContractItemDEX — items", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Item", "Material", "Text", "Target qty", "Net price", "Target amount", "Unit", "Plant", "GR", "Invoice"],
        [
            [
                item["purchaseContractItem"],
                item["material"],
                item["text"],
                item["targetQuantity"],
                item["netPrice"],
                item["targetAmount"],
                item["unit"],
                item["plant"],
                item["goodsReceiptIsExpected"],
                item["invoiceIsExpected"],
            ]
            for item in contract["items"]
        ],
    )
    add_para(
        doc,
        "InvoiceIsGoodsReceiptBased is blank on the contract items. On the release PO, "
        "item 00020 (TG20) is goods-receipt-based and item 00010 (TG11) is not.",
    )

    doc.add_heading("6. Release purchase order 4500002146", level=1)
    add_para(
        doc,
        "Standard PO (type NB), date 20261002, purchasing group 001, processing status 05, "
        "header total 315.00 USD. The contract link is on the PO item "
        "(PurchaseContract and PurchaseContractItem) and on contract history "
        "(ReleaseOrder and ReleaseOrderItem).",
    )
    add_para(doc, "C_PurchaseOrderDEX — header", size=11, bold=True, space_after=4)
    po_row = load_csv("C_PURCHASEORDERDEX.csv")[0]
    add_field_table(
        doc,
        nonempty_pairs(
            po_row,
            [
                "PURCHASEORDER",
                "PURCHASEORDERTYPE",
                "PURCHASINGDOCUMENTORIGIN",
                "CREATEDBYUSER",
                "CREATIONDATE",
                "PURCHASEORDERDATE",
                "LANGUAGE",
                "PURCHASINGPROCESSINGSTATUS",
                "COMPANYCODE",
                "PURCHASINGORGANIZATION",
                "PURCHASINGGROUP",
                "SUPPLIER",
                "PAYMENTTERMS",
                "DOCUMENTCURRENCY",
                "EXCHANGERATE",
                "PURGRELEASETIMETOTALAMOUNT",
                "LASTCHANGEDATETIME",
            ],
        ),
    )
    add_para(doc, "C_PurchaseOrderItemDEX — items", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["PO item", "Contract item", "Material", "Text", "Qty", "Net price", "Net amount", "Stor. loc", "GR-based IR"],
        [
            [
                item["purchaseOrderItem"],
                item["purchaseContractItem"],
                item["material"],
                item["text"],
                item["orderQuantity"],
                item["netPrice"],
                item["netAmount"],
                item["storageLocation"],
                item["invoiceIsGoodsReceiptBased"] or "(blank)",
            ]
            for item in po["items"]
        ],
    )
    add_para(doc, "C_PurchaseContractHistoryDEX — release rows for this contract", size=11, bold=True, space_after=4)
    history_rows = load_csv("C_PURCHASECONTRACTHISTORYDEX.csv")
    add_table(
        doc,
        ["Contract item", "Release order", "Release item", "Qty", "Net amount", "Date", "Plant"],
        [
            [
                row["PURCHASECONTRACTITEM"],
                row["RELEASEORDER"],
                row["RELEASEORDERITEM"],
                row["RELEASEORDERITEMORDERQUANTITY"],
                row["RELEASEORDERITEMNETAMOUNT"],
                row["RELEASEORDERDATE"],
                row["PLANT"],
            ]
            for row in history_rows
        ],
    )

    doc.add_heading("7. Goods receipts", level=1)
    add_para(
        doc,
        "C_PurchaseOrderHistoryDEX rows with PurchasingHistoryDocumentType 1 and "
        "PurchasingHistoryCategory E are goods receipts. The history document number is "
        "the material document. I_GoodsMovementDocumentDEX repeats that number with "
        "movement type 101, inventory transaction type WE, and goods-movement reference "
        "document type B (Goods movement for purchase order).",
    )
    add_para(doc, "C_PurchaseOrderHistoryDEX — type 1", size=11, bold=True, space_after=4)
    gr_hist = [
        row
        for row in correlation["purchaseOrderHistory"]
        if row["purchasingHistoryDocumentType"] == "1"
    ]
    add_table(
        doc,
        ["PO item", "Material doc", "Item", "Material", "Qty", "Amount", "Posting date", "Mvmt", "Contract item"],
        [
            [
                row["purchaseOrderItem"],
                row["purchasingHistoryDocument"],
                row["purchasingHistoryDocumentItem"],
                row["material"],
                row["quantity"],
                row["amount"],
                row["postingDate"],
                row["goodsMovementType"],
                row["purchaseContractItem"],
            ]
            for row in gr_hist
        ],
    )
    add_para(doc, "I_GoodsMovementDocumentDEX", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Material doc", "Year", "Item", "PO item", "Material", "Qty", "Unit", "Posting date", "Plant", "SLoc", "Ref. type"],
        [
            [
                row["materialDocument"],
                row["materialDocumentYear"],
                row["materialDocumentItem"],
                row["purchaseOrderItem"],
                row["material"],
                row["quantity"],
                row["unit"],
                row["postingDate"],
                row["plant"],
                row["storageLocation"],
                row["goodsMovementRefDocType"],
            ]
            for row in correlation["goodsReceipts"]
        ],
    )
    add_para(
        doc,
        "Each goods receipt posts on leading ledger 0L: debit stock account 0013600000, "
        "credit GR/IR account 0021120000. The accounting document is a different number "
        "from the material document. Join with ReferenceDocumentType MKPF and "
        "ReferenceDocument equal to the material document.",
    )
    add_table(
        doc,
        ["Material document", "Accounting document", "Posting date", "Amount USD", "PO item"],
        [
            ["5000002931", "5000000001", "20261002", "36.00", "00020"],
            ["5000002940", "5000000002", "20261005", "84.00", "00020"],
            ["5000002941", "5000000003", "20261005", "40.50", "00010"],
        ],
    )

    doc.add_heading("8. Supplier invoices", level=1)
    add_para(
        doc,
        "C_PurchaseOrderHistoryDEX rows with PurchasingHistoryDocumentType 2 and "
        "PurchasingHistoryCategory Q are invoice receipts. The history document number "
        "is the supplier invoice. Both invoices are posted (SupplierInvoiceStatus 5) and "
        "have a matching accounting document of type RE. Gross amount equals the sum of "
        "the item amounts. These journal lines have no separate tax account.",
    )
    add_para(doc, "C_SupplierInvoiceDEX — headers", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["Supplier invoice", "Fiscal year", "Party reference", "Invoicing party", "Posting date", "Gross USD", "Status"],
        [
            [
                row["supplierInvoice"],
                row["fiscalYear"],
                row["supplierInvoiceIdByInvcgParty"],
                row["invoicingParty"],
                row["postingDate"],
                row["grossAmount"],
                row["supplierInvoiceStatus"],
            ]
            for row in correlation["supplierInvoices"]
        ],
    )
    add_para(doc, "C_SupplierInvoiceItemDEX — items", size=11, bold=True, space_after=4)
    invoice_item_rows = []
    for invoice in correlation["supplierInvoices"]:
        for item in invoice["items"]:
            invoice_item_rows.append(
                [
                    invoice["supplierInvoice"],
                    item["supplierInvoiceItem"],
                    item["purchaseOrderItem"],
                    item["material"],
                    item["quantity"],
                    item["amount"],
                    item["referenceMaterialDocument"] or "(blank)",
                ]
            )
    add_table(
        doc,
        ["Supplier invoice", "Item", "PO item", "Material", "Qty", "Amount USD", "Material document"],
        invoice_item_rows,
    )
    add_para(
        doc,
        "PO item 00020 is goods-receipt-based, so invoice 5100001598 stores material "
        "documents 5000002931 and 5000002940 in PrmthbReferenceDocument. PO item 00010 "
        "is not goods-receipt-based, so invoice 5100001599 leaves that reference empty. "
        "Correlate item 00010 with PurchaseOrder and PurchaseOrderItem.",
    )
    add_para(doc, "C_PurchaseOrderHistoryDEX — type 2", size=11, bold=True, space_after=4)
    ir_hist = [
        row
        for row in correlation["purchaseOrderHistory"]
        if row["purchasingHistoryDocumentType"] == "2"
    ]
    add_table(
        doc,
        ["PO item", "Supplier invoice", "History item", "Material", "Qty", "Amount", "Posting date", "Ref. document", "Contract item"],
        [
            [
                row["purchaseOrderItem"],
                row["purchasingHistoryDocument"],
                row["purchasingHistoryDocumentItem"],
                row["material"],
                row["quantity"],
                row["amount"],
                row["postingDate"],
                row["referenceDocument"] or "(blank)",
                row["purchaseContractItem"],
            ]
            for row in ir_hist
        ],
    )

    doc.add_heading("9. Universal journal", level=1)
    add_para(
        doc,
        "I_GLAccountLineItemRawData, company code 1710, fiscal year 2026, source ledger 0L. "
        "The same documents also exist on ledger 2L; this sample keeps 0L. "
        "Join a supplier invoice with ReferenceDocumentType RMRP and ReferenceDocument "
        "equal to the supplier invoice. The vendor line (financial account type K) has "
        "no PurchasingDocument. The GR/IR lines on the same accounting document carry "
        "PurchasingDocument 4500002146 and the PO item. AssignmentReference is the PO "
        "and item concatenated, for example 450000214600020.",
    )
    add_para(doc, "G/L accounts used (I_GLAccount, chart YCOA, company 1710)", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["G/L account", "External", "Reconciliation", "Open item", "Role on this chain"],
        [
            ["0013600000", "13600000", "", "", "Stock, debited at goods receipt (financial account type M)"],
            ["0021120000", "21120000", "", "X", "GR/IR (financial account type S)"],
            ["0021100000", "21100000", "K", "", "Vendor payables (financial account type K)"],
        ],
    )
    add_para(
        doc,
        "I_GLAccount in this extract has no account-name column. The role comes from the "
        "posting and from ReconciliationAccountType / IsOpenItemManaged.",
        space_after=6,
    )
    add_para(doc, "Journal lines for this PO", size=11, bold=True, space_after=4)
    add_table(
        doc,
        ["FI document", "Type", "Line", "G/L account", "D/C", "Amount USD", "PO item", "Ref. type", "Reference document", "Assignment", "Item text"],
        [
            [
                line["accountingDocument"],
                line["accountingDocumentType"],
                line["ledgerGLLineItem"],
                line["glAccount"],
                line["debitCreditCode"],
                line["amountInCompanyCodeCurrency"],
                line["purchasingDocumentItem"] if line["purchasingDocument"] else "",
                line["referenceDocumentType"],
                line["referenceDocument"],
                line["assignmentReference"],
                line["documentItemText"],
            ]
            for line in correlation["journalLines"]
        ],
        font=7,
    )
    add_para(
        doc,
        "Goods-receipt documents 5000000001, 5000000002, and 5000000003 debit stock and "
        "credit GR/IR. Invoice documents 5100000000 and 5100000001 debit GR/IR and credit "
        "payables. GR/IR nets to 0.00 for the quantities that were received and invoiced. "
        "ClearingDate is 00000000 and ClearingAccountingDocument is blank on every line.",
    )

    doc.add_heading("10. Quantity and amount reconciliation", level=1)
    add_para(
        doc,
        "Invoiced quantity equals received quantity on both released items. The open "
        "amount is quantity not yet received.",
    )
    add_table(
        doc,
        ["Contract item", "PO item", "Material", "Target", "Released", "Ordered", "Order amt", "Received", "Recv amt", "Invoiced", "Inv amt", "Still to receive", "Not released"],
        [
            [
                row["purchaseContractItem"],
                row["purchaseOrderItem"],
                row["material"],
                row["contractTargetQuantity"],
                row["releasedQuantity"],
                row["orderQuantity"],
                row["orderAmount"],
                row["receivedQuantity"],
                row["receivedAmount"],
                row["invoicedQuantity"],
                row["invoicedAmount"],
                row["openReceiptQuantity"],
                row["unreleasedQuantity"],
            ]
            for row in correlation["quantityReconciliation"]
        ],
        font=7,
    )

    doc.add_heading("11. Payment", level=1)
    add_para(
        doc,
        f"Open vendor balance on G/L account {payment['glAccount']}: "
        f"{payment['openAmount']} {payment['currency']}. "
        "That is invoice 5100001598 (120.00) plus invoice 5100001599 (40.50). "
        f"GR/IR net amount is {payment['grirNetAmount']} {payment['currency']}.",
    )
    add_para(doc, payment["note"])
    add_para(
        doc,
        "Company 1710 in this extract has accounting document types RE, RV, WA, WE, and WL. "
        "A later payment would be a journal whose clearing fields point at accounting "
        "documents 5100000000 and 5100000001.",
    )

    doc.add_heading("12. How the documents are correlated", level=1)
    add_table(
        doc,
        ["From", "To", "Keys"],
        [
            ["C_PurchaseContractDEX", "C_PurchaseContractItemDEX", "PurchaseContract"],
            ["C_PurchaseContractItemDEX", "C_PurchaseContractHistoryDEX", "PurchaseContract + PurchaseContractItem. ReleaseOrder is the PO."],
            ["C_PurchaseContractItemDEX", "C_PurchaseOrderItemDEX", "PurchaseContract + PurchaseContractItem"],
            ["C_PurchaseOrderDEX", "C_PurchaseOrderItemDEX", "PurchaseOrder"],
            ["C_PurchaseOrderItemDEX", "C_PurchaseOrderHistoryDEX", "PurchaseOrder + PurchaseOrderItem. Type 1 = GR, type 2 = invoice."],
            ["History type 1", "I_GoodsMovementDocumentDEX", "PurchasingHistoryDocument = MaterialDocument. Movement type 101."],
            ["History type 2", "C_SupplierInvoiceDEX", "PurchasingHistoryDocument = SupplierInvoice, plus company code and fiscal year."],
            ["C_SupplierInvoiceItemDEX", "Material document", "PrmthbReferenceDocument when the PO item is GR-based."],
            ["Material document", "I_GLAccountLineItemRawData", "ReferenceDocumentType MKPF and ReferenceDocument = MaterialDocument."],
            ["Supplier invoice", "I_GLAccountLineItemRawData", "ReferenceDocumentType RMRP and ReferenceDocument = SupplierInvoice."],
            ["Journal GR/IR line", "PO item", "PurchasingDocument + PurchasingDocumentItem. AssignmentReference is the same pair with no separator."],
            ["Journal vendor line", "PO", "Same AccountingDocument as the GR/IR lines. The vendor line has no PurchasingDocument."],
            ["PO supplier", "I_BusinessPartner", "Supplier = BusinessPartner"],
            ["Journal line", "I_GLAccount", "GLAccount + CompanyCode"],
        ],
        font=8,
    )
    add_para(
        doc,
        "Material document 5000002931 posts to accounting document 5000000001. "
        "Supplier invoice 5100001598 posts to accounting document 5100000000. "
        "Join logistics documents to the journal on ReferenceDocument and ReferenceDocumentType.",
    )

    doc.add_heading("13. Extract calls for this reference", level=1)
    add_para(
        doc,
        "The files in this use case were written by ZEVO_CDS_EXPLORER_2_FILE in run "
        "20261006_094549. The same chain over OData is one ExtractCds call per entity. "
        "There is no $expand. Example for the contract header:",
    )
    add_para(
        doc,
        "GET /sap/opu/odata/sap/ZEVO_CDS_EXTRACT_SRV/ExtractCds"
        "?EntityName='C_PurchaseContractDEX'"
        "&Filter='PurchaseContract eq ''4600000041'''"
        "&Format='jsonrows'&Skip='0'&Top='10'&$format=json",
        size=9,
    )
    add_table(
        doc,
        ["Entity", "Filter"],
        [
            ["C_PurchaseContractItemDEX", "PurchaseContract eq '4600000041'"],
            ["C_PurchaseContractHistoryDEX", "PurchaseContract eq '4600000041'"],
            ["C_PurchaseOrderDEX", "PurchaseOrder eq '4500002146'"],
            ["C_PurchaseOrderItemDEX", "PurchaseOrder eq '4500002146'"],
            ["C_PurchaseOrderHistoryDEX", "PurchaseOrder eq '4500002146'"],
            ["I_GoodsMovementDocumentDEX", "PurchaseOrder eq '4500002146'"],
            ["C_SupplierInvoiceItemDEX", "PurchaseOrder eq '4500002146'"],
            ["C_SupplierInvoiceDEX", "SupplierInvoice eq '5100001598' or SupplierInvoice eq '5100001599'"],
            ["I_GLAccountLineItemRawData", "CompanyCode eq '1710' and SourceLedger eq '0L' and the five accounting documents in section 9"],
            ["I_BusinessPartner", "BusinessPartner eq '0001000559'"],
            ["I_BusinessPartnerSupplierDEX", "Supplier eq '0001000559'"],
            ["I_GLAccount", "CompanyCode eq '1710' and GLAccount eq '0013600000' or '0021100000' or '0021120000'"],
        ],
        font=8,
    )

    doc.add_heading("14. Sample-data files", level=1)
    add_para(
        doc,
        "CSV delimiter is semicolon. Headers are the CDS element names in uppercase. "
        "The tables in this document are the sample rows for contract 4600000041 and "
        "PO 4500002146. Regenerate the seed and docs/samples/c2p/correlation.json with "
        "python3 docs/samples/c2p/build_seed.py, then regenerate this document with "
        "python3 docs/samples/c2p/build_use_case_docx.py.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
