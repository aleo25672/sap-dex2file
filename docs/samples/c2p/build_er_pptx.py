#!/usr/bin/env python3
"""Build docs/CONTRACT_TO_PAYMENT_ER.pptx — entity diagram for contract to payment.

    python3 docs/samples/c2p/build_er_pptx.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[2] / "CONTRACT_TO_PAYMENT_ER.pptx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x0F, 0x6B, 0x6B)
AMBER = RGBColor(0x8A, 0x5A, 0x00)
PLUM = RGBColor(0x5C, 0x3D, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
PALE = RGBColor(0xF4, 0xF7, 0xFB)
LINE = RGBColor(0xC5, 0xD0, 0xDC)


def set_run(run, text, size, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(shape, lines, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT):
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        content, line_bold, line_color, line_size = line if isinstance(line, tuple) else (line, bold, color, size)
        set_run(paragraph.add_run(), content, line_size, line_bold, line_color)


def box(slide, left, top, width, height, fill, title, body):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.1)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.06)
    frame.margin_bottom = Inches(0.04)
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    set_run(frame.paragraphs[0].add_run(), title, 12, True, WHITE)
    for line in body:
        paragraph = frame.add_paragraph()
        paragraph.space_before = Pt(1)
        set_run(paragraph.add_run(), line, 10, False, WHITE)
    return shape


def label(slide, left, top, width, height, text, size=11, bold=True, color=NAVY, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(left, top, width, height)
    add_text(shape, [(text, bold, color, size)], align=align)
    return shape


def arrow(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = LINE
    shape.line.fill.background()
    return shape


def footer(slide, text):
    label(slide, Inches(0.4), Inches(7.15), Inches(12.5), Inches(0.28), text, size=11, bold=False, color=MUTED)


def main() -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = PALE
    background.line.fill.background()

    label(slide, Inches(0.4), Inches(0.22), Inches(10), Inches(0.38), "Contract to Payment", size=26, color=NAVY)
    label(
        slide,
        Inches(0.4),
        Inches(0.58),
        Inches(12),
        Inches(0.28),
        "Entity diagram  ·  contract 4600000041  ·  release PO 4500002146  ·  payment 1500000000",
        size=13,
        bold=False,
        color=MUTED,
    )

    columns = [
        (NAVY, "Contract", [
            ("BusinessPartner", ["BusinessPartner", "Supplier on the contract"]),
            ("PurchaseContract", ["PK PurchaseContract", "Supplier, CompanyCode"]),
            ("PurchaseContractItem", ["PK Contract + Item", "Material, target qty, price"]),
            ("ContractHistory", ["ReleaseOrder = PO", "ReleaseOrderItem"]),
        ]),
        (TEAL, "Purchase order", [
            ("PurchaseOrder", ["PK PurchaseOrder", "Supplier, CompanyCode"]),
            ("PurchaseOrderItem", ["PK PO + Item", "FK PurchaseContract + Item"]),
            ("PurchaseOrderHistory", ["Type 1 = goods receipt", "Type 2 = invoice"]),
        ]),
        (AMBER, "Receipt and invoice", [
            ("GoodsMovement", ["PK MaterialDocument", "Year + Item, movement 101"]),
            ("SupplierInvoice", ["PK Invoice + Year", "+ CompanyCode"]),
            ("SupplierInvoiceItem", ["FK PO + Item", "PrmthbReferenceDocument"]),
        ]),
        (PLUM, "Journal and payment", [
            ("JournalLine WE", ["ReferenceDocumentType MKPF", "ReferenceDocument = mat. doc"]),
            ("JournalLine RE", ["ReferenceDocumentType RMRP", "ReferenceDocument = invoice"]),
            ("JournalLine KZ", ["Payment document", "ClearingAccountingDocument"]),
            ("GLAccount", ["PK GLAccount + CompanyCode", "Stock, GR/IR, vendor, bank"]),
        ]),
    ]

    lefts = [0.35, 3.62, 6.89, 10.16]
    width = 2.95
    for index, (color, heading, cards) in enumerate(columns):
        left = Inches(lefts[index])
        label(slide, left, Inches(0.98), Inches(width), Inches(0.28), heading, size=13, color=color)
        top = 1.32
        gap = 0.08
        height = 1.22 if len(cards) == 4 else 1.55
        for title, body in cards:
            box(slide, left, Inches(top), Inches(width), Inches(height), color, title, body)
            top += height + gap
        if index < 3:
            arrow(slide, Inches(lefts[index] + width + 0.02), Inches(3.35), Inches(0.28), Inches(0.18))

    footer(slide, "A payment is a journal document of type KZ. It is not a separate CDS entity.")

    slide = presentation.slides.add_slide(blank)
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = WHITE
    background.line.fill.background()
    label(slide, Inches(0.4), Inches(0.22), Inches(12), Inches(0.36), "How the entities connect", size=26, color=NAVY)
    label(
        slide,
        Inches(0.4),
        Inches(0.58),
        Inches(12),
        Inches(0.26),
        "Worked sample 4600000041 / 4500002146. The FI document number is different from the logistics number.",
        size=13,
        bold=False,
        color=MUTED,
    )

    rows = [
        ["Contract item → PO item", "PurchaseContract + PurchaseContractItem", "4600000041 / 00020 and 00030"],
        ["Contract history → PO", "ReleaseOrder + ReleaseOrderItem", "4500002146"],
        ["History type 1 → goods receipt", "PurchasingHistoryDocument = MaterialDocument", "5000002931, 5000002940, 5000002941"],
        ["History type 2 → invoice", "PurchasingHistoryDocument = SupplierInvoice", "5100001598, 5100001599"],
        ["Invoice item → goods receipt", "PrmthbReferenceDocument", "Filled when the item is GR-based"],
        ["Material document → journal", "ReferenceDocumentType MKPF", "FI 5000000001, 5000000002, 5000000003"],
        ["Invoice → journal", "ReferenceDocumentType RMRP", "FI 5100000000, 5100000001"],
        ["Invoice journal → payment", "ClearingAccountingDocument", "KZ 1500000000, 160.50 USD"],
        ["Supplier", "Supplier = BusinessPartner", "0001000559"],
        ["Journal → G/L", "GLAccount + CompanyCode", "13600000, 21120000, 21100000, 11002000"],
    ]

    table_shape = slide.shapes.add_table(1 + len(rows), 3, Inches(0.4), Inches(1.05), Inches(12.5), Inches(5.6))
    table = table_shape.table
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(4.7)
    table.columns[2].width = Inches(4.3)
    headers = ["Relationship", "Key", "In this sample"]
    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = ""
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_run(cell.text_frame.paragraphs[0].add_run(), text, 12, True, WHITE)
    for r_index, row in enumerate(rows):
        for col, text in enumerate(row):
            cell = table.cell(r_index + 1, col)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALE if r_index % 2 == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            set_run(cell.text_frame.paragraphs[0].add_run(), text, 12, False, INK)

    footer(slide, "ZEVO_C2P navigates the purchasing entities. The journal stays in I_GLAccountLineItemRawData.")

    presentation.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
